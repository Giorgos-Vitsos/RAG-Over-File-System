from pathlib import Path
import docx
from docx.table import Table
from .tabular import format_rows,fill_merged
from docx.document import Document as DocxDocument
import io
import zipfile
import pptx
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.presentation import Presentation as PptxPresentation
from striprtf.striprtf import rtf_to_text
from .text import read_text
from docx.oxml.ns import qn
from .text import clean_lines,with_address
from pptx.oxml.ns import qn as pptx_qn
import re
from docx.oxml import parse_xml

DOCX_MAIN = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"#type of normal docx file
WORD_VARIANTS=[#types of doc variants
    "application/vnd.ms-word.document.macroEnabled.main+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
    "application/vnd.ms-word.template.macroEnabledTemplate.main+xml"
]

PPTX_MAIN="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"#type of normal pptx file
SLIDE_VARIANTS=[#types of ppt variants
    "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
    "application/vnd.ms-powerpoint.template.macroEnabled.main+xml",
    "application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml",
    "application/vnd.ms-powerpoint.slideshow.macroEnabled.main+xml"
]

MC="{http://schemas.openxmlformats.org/markup-compatibility/2006}"
WP_DOCPR="{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"

DOCX_SKIP={#tags without text of the document
    qn("w:del"),qn("w:moveFrom"),qn("w:delText"),qn("w:instrText"),
    qn("w:pPr"),qn("w:rPr"),qn("w:tcPr"),qn("w:sdtPr"),qn("w:sectPr"),
    qn("w:footnoteReference"),qn("w:endnoteReference"),qn("w:commentReference")
}

FALSE_VALUES=("0","false","off")

FIELD_LINK=re.compile(r"\s*hyperlink\b(.*)",re.I|re.S)#the word HYPERLINK and what follows it
FIELD_SWITCH=re.compile(r'\\[lot]\s+(?:"[^"]*"|\S+)|\\[mn]\b',re.I)#more arguments of a link field
FIELD_ADDRESS=re.compile(r'\s*(?:"([^"]*)"|(\S+))')
CHART="{http://schemas.openxmlformats.org/drawingml/2006/chart}"
DRAWING="{http://schemas.openxmlformats.org/drawingml/2006/main}"
RELATION="{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
DGM="{http://schemas.openxmlformats.org/drawingml/2006/diagram}"#tags of a smartart

def heading_level(style_name:str)->int:
    if style_name=="Title":
        return 1
    elif style_name.startswith("Heading"):
        try:
            return int(style_name.split()[-1])
        except ValueError:
            return 0
    else:
        return 0

def docx_to_text(doc: DocxDocument)->str:
    headings={}
    for style in doc.styles:
        headings[style.style_id]=heading_level(str(style.name))
    return "\n\n".join(docx_blocks(doc.element.body,doc,headings))

def docx_blocks(node,doc:DocxDocument,headings:dict[str,int] | None)->list[str]:
    blocks=[]
    for child in node:
        blocks.extend(docx_block(child,doc,headings))
    return blocks

def docx_block(node,doc:DocxDocument,headings:dict[str,int] | None)->list[str]:
    tag=node.tag
    if tag in DOCX_SKIP:
        return []
    if tag==qn("w:p"):
        return docx_paragraph(node,doc,headings)
    if tag==qn("w:tbl"):
        text=docx_table(Table(node,doc._body),doc)
        if text:
            return [text]
        return []
    if tag==qn("w:sdt") and is_toc(node):#the table of contents repeats data so we skip it
        return []
    if tag==MC+"AlternateContent":
        return docx_blocks(node[0],doc,headings)#we read only node[0] because next node repeats the data
    return docx_blocks(node,doc,headings)#in wrappers and groups we search the children

def is_toc(sdt)->bool:
    properties=sdt.find(qn("w:sdtPr"))
    if properties is None:
        return False
    for gallery in properties.iter(qn("w:docPartGallery")):
        if gallery.get(qn("w:val"))=="Table of Contents":
            return True
    return False

def docx_paragraph(p,doc:DocxDocument,headings:dict[str,int] | None)->list[str]:
    shapes=[]#pictures and text boxes
    text=clean_lines(docx_inline(p,shapes,doc))
    blocks=[]
    for shape in shapes:
        blocks.extend(docx_shape(shape,doc,headings))
    if not text:
        return blocks
    level=0
    style=p.find(qn("w:pPr")+"/"+qn("w:pStyle"))
    if headings is not None and style is not None:#we add the headings
        level=headings.get(style.get(qn("w:val")),0)
    if level>0:
        text="#"*level+" "+text
    blocks.append(text)
    return blocks

def docx_inline(node,shapes:list,doc:DocxDocument)->str:
    parts=[]
    fields=[]
    for child in node:
        tag=child.tag
        if tag==qn("w:r"):#link seperated in different runs
            mark=child.find(qn("w:fldChar"))#all seperation marks
            if mark is not None:
                kind=mark.get(qn("w:fldCharType"))
                if kind=="begin":
                    fields.append({"code":"","result":None})
                elif kind=="separate" and fields:
                    fields[-1]["result"]=[]
                elif kind=="end" and fields:
                    field=fields.pop()
                    add_text(with_address("".join(field["result"] or []),field_address(field["code"])),parts,fields)
                continue
            codes=child.findall(qn("w:instrText"))
            if codes:
                if fields:
                    for code in codes:
                        fields[-1]["code"]+=code.text or ""
                continue
        if tag in DOCX_SKIP:
            continue
        if tag==qn("w:t"):
            add_text(child.text or "",parts,fields)
        elif tag==qn("w:tab"):
            add_text("\t",parts,fields)
        elif tag in (qn("w:br"),qn("w:cr")):
            add_text("\n",parts,fields)
        elif tag==qn("w:noBreakHyphen"):
            add_text("-",parts,fields)
        elif tag in (qn("w:drawing"),qn("w:pict")):
            shapes.append(child)
        elif tag==qn("w:fldSimple"):#different way of reading links
            text=docx_inline(child,shapes,doc)
            code=child.get(qn("w:instr")) or ""
            address=field_address(code)
            add_text(with_address(text,address),parts,fields)
        elif tag==qn("w:hyperlink"):
            add_text(with_address(docx_inline(child,shapes,doc),docx_link(child,doc)),parts,fields)
        elif tag==MC+"AlternateContent":
            add_text(docx_inline(child[0],shapes,doc),parts,fields)#we read only node[0] because next node repeats the data
        else:
            add_text(docx_inline(child,shapes,doc),parts,fields)
    while fields:
        field=fields.pop()
        add_text("".join(field["result"] or []),parts,fields)
    return "".join(parts)

def docx_link(link,doc:DocxDocument)->str:
    relation=doc.part.rels.get(link.get(qn("r:id")))#the address is kept here
    if relation is None or not relation.is_external:
        return ""
    return relation.target_ref

def docx_shape(shape,doc:DocxDocument,headings:dict[str,int] | None)->list[str]:
    blocks=[]
    for properties in outside_boxes(shape,WP_DOCPR):
        alt=(properties.get("descr") or properties.get("title") or "").strip()#we get either the desc or title or empty
        if alt:
            blocks.append(f"[Image: {alt}]")
    for chart in outside_boxes(shape,CHART+"chart"):
        part=doc.part.related_parts.get(chart.get(RELATION+"id",""))
        if part is not None:
            blocks.append(chart_text(parse_xml(part.blob)))
    for box in outside_boxes(shape,qn("w:txbxContent")):#it holds text and no figures
        blocks.extend(docx_blocks(box,doc,headings))
    for ids in outside_boxes(shape,DGM+"relIds"):
        part=doc.part.related_parts.get(ids.get(RELATION+"dm",""))
        if part is not None:
            text=diagram_text(parse_xml(part.blob))
            if text:
                blocks.append(text)
    return blocks

def outside_boxes(node,tag:str)->list:
    found=[]
    for child in node:
        if child.tag==tag:
            found.append(child)
        elif child.tag not in (qn("w:txbxContent"),MC+"Fallback"):
            found.extend(outside_boxes(child,tag))
    return found

def docx_table(table:Table,doc:DocxDocument)->str:
    rows=[]
    for row in table.rows:
        cells=[]
        for cell in row.cells:
            cells.append("\n".join(docx_blocks(cell._tc,doc,None)))
        rows.append(cells)
    return format_rows(rows,docx_header_rows(table))

def docx_header_rows(table:Table)->int:
    count=0
    for row in table.rows:#only the rows at the top can be the header
        properties=row._tr.trPr
        header=None
        if properties is not None:
            header=properties.find(qn("w:tblHeader"))
        if header is None or header.get(qn("w:val")) in FALSE_VALUES:#if no more header rows
            break
        count+=1
    return count

def read_docx(path: Path)->str:
    doc=docx.Document(str(path))
    return docx_to_text(doc)

#variants of some files can still be parsed as the standard file, for example in docm we dont need the macros so we parse it as docx
def variants_to_main(path: Path,main: str,variants: list[str])->io.BytesIO:
    buffer=io.BytesIO()
    with zipfile.ZipFile(path) as src,zipfile.ZipFile(buffer, "w") as dst:#we open the zip file and a virtual one
        for item in src.infolist():
            data=src.read(item.filename)
            if item.filename=="[Content_Types].xml":#if the file is the one which holds the type of the variant
                for variant in variants:
                    data=data.replace(variant.encode(),main.encode())#we change it to the standard of that file type
            dst.writestr(item,data)#we copy everything to the virtual file
    return buffer

def read_word_variants(path: Path)->str:
    doc=docx.Document(variants_to_main(path,DOCX_MAIN,WORD_VARIANTS))
    return docx_to_text(doc)

def read_slide_variants(path: Path)->str:
    ppt=pptx.Presentation(variants_to_main(path,PPTX_MAIN,SLIDE_VARIANTS))
    return pptx_to_text(ppt)

def pptx_link(run,part)->str:
    properties=run.find(pptx_qn("a:rPr"))
    if properties is None:
        return ""
    link=properties.find(pptx_qn("a:hlinkClick"))
    if link is None:
        return ""
    relation=part.rels.get(link.get(pptx_qn("r:id")))
    if relation is None or not relation.is_external:#jump inside the pptx
        return ""
    return relation.target_ref
    
def pptx_paragraph_text(paragraph)->str:#it builds the paragraph by fixing \x0b characters and finding links
    pieces=[]
    for child in paragraph._p:
        if child.tag==pptx_qn("a:br"):
            pieces.append(["\n",""])
            continue
        if child.tag not in(pptx_qn("a:r"),pptx_qn("a:fld")):#no text
            continue
        text=""
        for t in child.iter(pptx_qn("a:t")):
            if t.text is not None:
                text=text+t.text
        address=pptx_link(child,paragraph.part)
        if pieces and address and pieces[-1][1]==address:
            pieces[-1][0]+=text
        else:
            pieces.append([text,address])
    parts=[]
    for text,address in pieces:
        parts.append(with_address(text,address))
    return "".join(parts)


def pptx_to_text(ppt: PptxPresentation)->str:
    output=[]
    for i,slide in enumerate(ppt.slides):
        block=[]
        title=slide.shapes.title
        for shape in slide.shapes:
            if title is not None and shape.shape_id==title.shape_id:#if we have a title we dont want to write it twice
                continue
            block.extend(format_shape(shape))
        header=f"# Slide {i+1}"
        if title is not None:#we build the header with the title if it exists
            title_lines=[]
            for line in title.text_frame.text.replace("\x0b","\n").split("\n"):#title with multiple lines turns to single line
                title_txt=line.strip()
                if title_txt:
                    title_lines.append(title_txt)
            if title_lines:
                header=header+": "+" ".join(title_lines)        
        output.append(header)
        output.extend(block)
        if slide.has_notes_slide:#we add notes if they exist
            notes=slide.notes_slide.notes_text_frame
            if notes is not None:
                notes_txt=pptx_frame_text(notes)
                if notes_txt:
                    output.append(f"Notes: {notes_txt}")     
    return "\n\n".join(output)

def pptx_frame_text(frame)->str:
    paragraphs=[]
    for paragraph in frame.paragraphs:
        text=clean_lines(pptx_paragraph_text(paragraph))
        if text:
            paragraphs.append(text)
    return "\n".join(paragraphs)

def read_pptx(path: Path)->str:
    ppt=pptx.Presentation(str(path))
    return pptx_to_text(ppt)

def format_shape(shape:BaseShape)->list[str]:
    output=[]
    if isinstance(shape, GroupShape):
        for sub_shape in shape.shapes:
            output.extend(format_shape(sub_shape))
    elif isinstance(shape, Shape):
        text=pptx_frame_text(shape.text_frame)
        if text:
            output.append(text)
    elif isinstance(shape,GraphicFrame):
        if shape.has_chart:
            output.append(chart_text(shape.chart._chartSpace))
        elif shape.has_table:
            table=shape.table
            merges=[]
            for r in range(len(table.rows)):
                for c in range(len(table.columns)):
                    cell=table.cell(r,c)
                    if cell.is_merge_origin:#the top left cell of a merge knows its size
                        merges.append((r,c,r+cell.span_height-1,c+cell.span_width-1))
            rows=[]
            for row in table.rows:
                cells=[]
                for cell in row.cells:
                    cells.append(pptx_frame_text(cell.text_frame))
                rows.append(cells)
            text=format_rows(fill_merged(rows,merges)).strip()
            if text:
                output.append(text)
        else:
            ids=next(shape._element.iter(DGM+"relIds"),None)#we search in all depths for this
            if ids is not None:#we found smartart
                data=shape.part.related_part(ids.get(RELATION+"dm",""))
                text=diagram_text(parse_xml(data.blob))
                if text:
                    output.append(text)  
    for alt in shape._element.xpath("./*/p:cNvPr/@descr"):
        if alt.strip():
            output.append(f"[Image: {alt}]")
    return output

def read_rtf(path: Path)->str:
    text=read_text(path).strip()
    if not text.startswith(r"{\rtf1"):#if it doesnt start with that then its not a real .rtf file
        raise ValueError("Not an .rtf file")
    final=rtf_to_text(text).strip()
    return final

def field_address(code:str)->str:
    link=FIELD_LINK.match(code)
    if link is None:
        return ""
    rest=FIELD_SWITCH.sub(" ",link.group(1))
    address=FIELD_ADDRESS.match(rest)
    if address is None:#no address found
        return ""
    return address.group(1) or address.group(2) or ""

def add_text(text:str,parts:list[str],fields:list[dict])->None:#puts text in the correct pos
    if not fields:#outside of every field
        parts.append(text)
    elif fields[-1]["result"] is not None:#inside the result of a field
        fields[-1]["result"].append(text)

def chart_text(space)->str:
    title=space.find(CHART+"chart/"+CHART+"title")
    title_parts=[]
    if title is not None:
        for paragraph in title.iter(DRAWING+"p"):
            pieces=[]
            for piece in paragraph.iter(DRAWING+"t"):
                pieces.append(piece.text or "")
            title_parts.append("".join(pieces))
    name=" ".join(" ".join(title_parts).split())
    if name:#if we have a title for the chart
        lines=[f"[Chart: {name}]"]
    else:
        lines=["[Chart]"]
    categories={}
    values={}
    for series in space.iter(CHART+"ser"):
        series_name=""
        tx=series.find(CHART+"tx")
        if tx is not None:#if there is a name
            names=[]
            for v in tx.iter(CHART+"v"):
                names.append(v.text or "")
            series_name=" ".join(names)#we put the real name if it exists
        category=series.find(CHART+"cat")
        if category is None:#scatter charts have different category
            category=series.find(CHART+"xVal")
        value=series.find(CHART+"val")
        if value is None:#same as before
            value=series.find(CHART+"yVal")
        if category is not None:
            for point in category.iter(CHART+"pt"):
                categories.setdefault(int(point.get("idx") or 0),point.findtext(CHART+"v") or "")
        if value is not None:
            for point in value.iter(CHART+"pt"):
                number=point.findtext(CHART+"v") or ""
                values.setdefault(int(point.get("idx") or 0),[]).append(f"{series_name} {number}".strip())
    for index in sorted(categories):
        if index in values:
            lines.append(f"{categories[index]}: "+"; ".join(values[index]))
    return "\n".join(lines)

def diagram_text(model)->str:
    kinds={}
    texts={}
    for point in model.iter(DGM+"pt"):
        point_id=point.get("modelId") or ""
        kinds[point_id]=point.get("type") or "node"#a point without a type is a normal item
        body=point.find(DGM+"t")
        if body is not None:
            paragraphs=[]
            for paragraph in body.iter(DRAWING+"p"):
                pieces=[]
                for piece in paragraph.iter(DRAWING+"t"):
                    pieces.append(piece.text or "")
                paragraphs.append("".join(pieces))
            texts[point_id]=" ".join(" ".join(paragraphs).split())
    children={}
    for link in model.iter(DGM+"cxn"):
        if (link.get("type") or "parOf")!="parOf":#the other ones are for drawings
            continue
        children.setdefault(link.get("srcId") or "",[]).append((int(link.get("srcOrd") or 0),link.get("destId") or ""))
    lines=[]
    for point_id,kind in kinds.items():
        if kind=="doc":#the root
            diagram_lines(point_id,0,kinds,texts,children,lines)
    if not lines:
        return ""
    return "\n".join(["[Diagram]"]+lines)

def diagram_lines(parent:str,depth:int,kinds:dict,texts:dict,children:dict,lines:list[str])->None:#crates a markdown like hirarchy
    for _,child in sorted(children.get(parent,[])):
        if kinds.get(child) not in ("node","asst"):#points only for drawing
            continue
        text=texts.get(child,"")
        if text:
            lines.append("  "*depth+"- "+text)
            diagram_lines(child,depth+1,kinds,texts,children,lines)
        else:#if empty the children go there
            diagram_lines(child,depth,kinds,texts,children,lines)