import datetime
import io
import itertools
import re
import zipfile

import docx
import pytest
from odf.dc import Creator, Date as DcDate, Title as DcTitle
from odf.draw import (
    A as DrawA, CustomShape, Frame, G, Image as DrawImage, Object as DrawObject, Page, PageThumbnail, TextBox,
)
from odf.namespaces import TEXTNS, XMLNS
from odf.office import Annotation, AnnotationEnd, ChangeInfo
from odf.opendocument import OpenDocumentPresentation, OpenDocumentSpreadsheet, OpenDocumentText
from odf.presentation import FooterDecl, Notes
from odf.style import DrawingPageProperties, Footer, Header, MasterPage, PageLayout, Style
from odf.svg import Desc, Title as SvgTitle
from odf.table import (
    CoveredTableCell, Table as OdfTable, TableCell, TableColumn, TableHeaderRows, TableRow, TableRowGroup,
)
from odf.text import (
    A, AlphabeticalIndex, Bibliography, Change, ChangedRegion, ChangeEnd, ChangeStart, Date as TextDate, Deletion, H,
    IndexBody, IndexTitle, Insertion, LineBreak as OdfLineBreak, List, ListHeader, ListItem, Note, NoteBody,
    NoteCitation, NumberedParagraph, P as OdfP, S as OdfS, Section, Sequence, SoftPageBreak, Span, Tab,
    TableOfContent, TrackedChanges,
)
from PIL import Image

from rag_fs.ingest.parsers import PARSERS
from rag_fs.ingest.parsers.opendocument import read_odf, read_ods
from test_parsers import EXPECTED_DOCX, EXPECTED_PPTX, EXPECTED_TWO_SHEETS, INCOME, built, make_ods, parse_one


# ---------------- .ods ----------------

# the same for an .ods file
def test_read_ods_all_sheets(tmp_path):
    f = tmp_path / "book.ods"
    rows = [["Ημερομηνία", "Κατηγορία", "Ποσό"], ["2024-03-01", "Ρεύμα", 54]]
    make_ods(f, [("Έξοδα", rows), ("Κενό", []), ("Έσοδα", INCOME)])

    assert read_ods(f) == EXPECTED_TWO_SHEETS


# ods cells and rows written the way libreoffice writes them
def ods_cell(*paragraphs, **attrs):
    cell = TableCell(**attrs)
    for p in paragraphs:
        cell.addElement(OdfP(text=p) if isinstance(p, str) else p)
    return cell


def ods_row(*cells, **attrs):
    row = TableRow(**attrs)
    for cell in cells:
        row.addElement(cell)
    return row


# a paragraph with two spaces, a line break and more text
def spaced_paragraph():
    p = OdfP(text="β")
    p.addElement(OdfS(c=2))
    p.addText("γ")
    p.addElement(OdfLineBreak())
    p.addText("δ")
    return p


def commented_cell():
    note = Annotation()
    note.addElement(OdfP(text="σχόλιο"))
    cell = ods_cell(valuetype="string")
    cell.addElement(note)
    cell.addElement(OdfP(text="Μαρία"))
    return cell


# a cell with a bullet list in it
def list_cell(*items):
    bullet_list = List()
    for item in items:
        list_item = ListItem()
        list_item.addElement(OdfP(text=item))
        bullet_list.addElement(list_item)
    return ods_cell(bullet_list, valuetype="string")


S_ = lambda text: ods_cell(text, valuetype="string")  # noqa: E731
names = lambda: ods_row(S_("Όνομα"), S_("Γραπτό"), S_("Προφορικό"))  # noqa: E731
ODS_CASES = {
    "empty tail like libreoffice": ([
        names(),
        ods_row(S_("Μαρία"), ods_cell("9", valuetype="float", value="9"), S_("8"), ods_cell(numbercolumnsrepeated=1021)),
        ods_row(ods_cell(numbercolumnsrepeated=1024), numberrowsrepeated=1048574),
    ], "Όνομα: Μαρία; Γραπτό: 9; Προφορικό: 8"),
    "value repeated across columns": ([
        names(), ods_row(S_("Μαρία"), ods_cell("9", valuetype="float", value="9", numbercolumnsrepeated=2)),
    ], "Όνομα: Μαρία; Γραπτό: 9; Προφορικό: 9"),
    "row repeated": ([
        names(), ods_row(S_("Μαρία"), S_("9"), S_("8"), numberrowsrepeated=2),
    ], "Όνομα: Μαρία; Γραπτό: 9; Προφορικό: 8\nΌνομα: Μαρία; Γραπτό: 9; Προφορικό: 8"),
    "empty rows between": ([
        names(), ods_row(S_("Μαρία"), S_("9"), S_("8")),
        ods_row(ods_cell(numbercolumnsrepeated=3), numberrowsrepeated=4),
        ods_row(S_("Νίκος"), S_("7"), S_("6")),
    ], "Όνομα: Μαρία; Γραπτό: 9; Προφορικό: 8\nΌνομα: Νίκος; Γραπτό: 7; Προφορικό: 6"),
    "cells covered by a merge, repeated": ([
        ods_row(S_("Όνομα"), S_("Γραπτό"), S_("Προφορικό"), S_("Εργασία")),
        ods_row(S_("Μαρία"), ods_cell("απαλλαγή", valuetype="string", numbercolumnsspanned=3),
                CoveredTableCell(numbercolumnsrepeated=2)),
    ], "Όνομα: Μαρία; Γραπτό: απαλλαγή; Προφορικό: απαλλαγή; Εργασία: απαλλαγή"),
    "date with time and numbers shown differently": ([
        ods_row(S_("Ημερομηνία"), S_("Ποσό"), S_("Ποσοστό"), S_("Τιμή")),
        ods_row(ods_cell("01/03/24 14:30", valuetype="date", datevalue="2024-03-01T14:30:00"),
                ods_cell("54,50", valuetype="float", value="54.5"),
                ods_cell("25%", valuetype="percentage", value="0.25"),
                ods_cell("12,00 €", valuetype="currency", currency="EUR", value="12")),
    ], "Ημερομηνία: 2024-03-01 14:30:00; Ποσό: 54.5; Ποσοστό: 0.25; Τιμή: 12"),
    "paragraphs, spaces and line breaks in a cell": ([
        ods_row(S_("Όνομα"), S_("Σημειώσεις")),
        ods_row(S_("Μαρία"), ods_cell("α", spaced_paragraph(), valuetype="string")),
    ], "Όνομα: Μαρία; Σημειώσεις: α / β  γ / δ"),
    "merge in the empty area below the table": ([
        names(), ods_row(S_("Μαρία"), S_("9"), S_("8")),
        ods_row(ods_cell(numbercolumnsrepeated=3), numberrowsrepeated=5),
        ods_row(ods_cell(numbercolumnsspanned=2), CoveredTableCell(), ods_cell()),
    ], "Όνομα: Μαρία; Γραπτό: 9; Προφορικό: 8"),
    "merge after empty rows and an empty cell": ([
        ods_row(S_("Όνομα"), S_("Σχόλιο"), S_("Γραπτό"), S_("Προφορικό")),
        ods_row(S_("Μαρία"), ods_cell(), S_("9"), S_("8")),
        ods_row(ods_cell(numbercolumnsrepeated=4), numberrowsrepeated=2),
        ods_row(S_("Νίκος"), ods_cell(), ods_cell("απαλλαγή", valuetype="string", numbercolumnsspanned=2),
                CoveredTableCell()),
    ], "Όνομα: Μαρία; Γραπτό: 9; Προφορικό: 8\nΌνομα: Νίκος; Γραπτό: απαλλαγή; Προφορικό: απαλλαγή\nΣχόλιο"),
    "a comment is not the cell text": ([
        ods_row(S_("Όνομα"), S_("Βαθμός")), ods_row(commented_cell(), S_("9")),
    ], "Όνομα: Μαρία; Βαθμός: 9"),
    "spaces and new lines in the xml of a cell count as one space": ([
        ods_row(S_("Όνομα"), S_("Επώνυμο")), ods_row(S_("Μαρία"), S_("Παπαδάκη\n      Λαμπράκη")),
    ], "Όνομα: Μαρία; Επώνυμο: Παπαδάκη Λαμπράκη"),
    "list in a cell": ([
        ods_row(S_("Όνομα"), S_("Μαθήματα")),
        ods_row(S_("Μαρία"), list_cell("ΗΥ100", "ΗΥ120")),
    ], "Όνομα: Μαρία; Μαθήματα: ΗΥ100 / ΗΥ120"),
}


# ods files as libreoffice writes them
@pytest.mark.parametrize("rows, expected", ODS_CASES.values(), ids=ODS_CASES.keys())
def test_read_ods_cells_and_rows(tmp_path, rows, expected):
    f = tmp_path / "book.ods"
    write_ods(f, rows)

    assert read_ods(f) == "# Sheet: Φύλλο\n" + expected


# writes an ods with one sheet; the first header_rows rows go in table:table-header-rows, the rest in a row group
def write_ods(path, rows, header_rows=0, group=False):
    doc = OpenDocumentSpreadsheet()
    table = OdfTable(name="Φύλλο")
    header = TableHeaderRows()
    body = TableRowGroup() if group else table
    if header_rows:
        table.addElement(header)
    if group:
        table.addElement(body)
    for i, row in enumerate(rows):
        (header if i < header_rows else body).addElement(row)
    doc.spreadsheet.addElement(table)  # pyright: ignore[reportAttributeAccessIssue] (odfpy adds it at runtime)
    doc.save(str(path))


# the header the file marks is used even when it is only numbers; rows inside groups are read
@pytest.mark.parametrize("header_rows, group, expected", [
    (0, False, "2023 | 2024\n120 | 340"),
    pytest.param(1, False, "2023: 120; 2024: 340", id="header marked by the file"),
    pytest.param(1, True, "2023: 120; 2024: 340", id="header marked, data in a row group"),
])
def test_read_ods_header_rows(tmp_path, header_rows, group, expected):
    f = tmp_path / "book.ods"
    years = lambda *v: ods_row(*[ods_cell(x, valuetype="float", value=x) for x in v])  # noqa: E731
    write_ods(f, [years("2023", "2024"), years("120", "340")], header_rows, group)

    assert read_ods(f) == "# Sheet: Φύλλο\n" + expected


# some programs put spaces and new lines between the tags of content.xml
def test_read_ods_pretty_printed_xml(tmp_path):
    plain = tmp_path / "plain.ods"
    write_ods(plain, [names(), ods_row(S_("Μαρία"), S_("9"), S_("8"))])
    pretty = tmp_path / "pretty.ods"
    with zipfile.ZipFile(plain) as zin, zipfile.ZipFile(pretty, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "content.xml":
                data = re.sub(rb">\s*<(?!/text:p)", b">\n   <", data)
            zout.writestr(item, data)

    assert read_ods(pretty) == read_ods(plain) == "# Sheet: Φύλλο\nΌνομα: Μαρία; Γραπτό: 9; Προφορικό: 8"


# a text document renamed to .ods is not a spreadsheet
def test_read_ods_refuses_a_text_document(tmp_path):
    doc = OpenDocumentText()
    doc.text.addElement(OdfP(text="κείμενο"))  # pyright: ignore[reportAttributeAccessIssue] (odfpy adds it at runtime)
    f = tmp_path / "text.ods"
    doc.save(str(f))

    with pytest.raises(ValueError):
        read_ods(f)


# a real date cell in ods should look like in xlsx, without 00:00:00
def test_read_ods_real_date_cells(tmp_path):
    f = tmp_path / "dates.ods"
    make_ods(f, [("Έξοδα", [["Ημερομηνία", "Ποσό"], [datetime.datetime(2024, 3, 1), 54]])])

    assert read_ods(f) == "# Sheet: Έξοδα\nΗμερομηνία: 2024-03-01; Ποσό: 54"


# ---------------- .odt ----------------

# a small png, every test document carries it like libreoffice does
def png_bytes():
    buffer = io.BytesIO()
    Image.new("RGB", (60, 30), "white").save(buffer, "PNG")
    return buffer.getvalue()


PNG = png_bytes()
PICTURE = "Pictures/img.png"


# a paragraph (or heading with cls=H) made of strings and odf elements
def para(*parts, cls=OdfP, **attrs):
    p = cls(**attrs)
    for part in parts:
        if isinstance(part, str):
            p.addText(part)
        else:
            p.addElement(part)
    return p


def heading(level, *parts):
    return para(*parts, cls=H, outlinelevel=level)


# puts the elements inside a parent element
def inside(parent, *children):
    for child in children:
        parent.addElement(child)
    return parent


# an image with its alt text (svg:title and svg:desc, like libreoffice writes them)
def image(title=None, desc=None):
    frame = Frame(width="2cm", height="1cm", anchortype="as-char")
    frame.addElement(DrawImage(href=PICTURE, type="simple", show="embed", actuate="onLoad"))
    if title is not None:
        frame.addElement(SvgTitle(text=title))
    if desc is not None:
        frame.addElement(Desc(text=desc))
    return frame


# a frame with a text box, like insert > text box in writer
def textbox(*paragraphs, anchortype="paragraph"):
    return inside(Frame(width="5cm", height="2cm", anchortype=anchortype), inside(TextBox(), *paragraphs))


# a chart (an object in another file of the zip); libreoffice also saves a picture of it to show
def chart(replacement_title=None):
    frame = inside(Frame(width="5cm", height="3cm", anchortype="paragraph"), DrawObject(href="./Object 1"))
    if replacement_title is not None:
        frame.addElement(DrawImage(href="./ObjectReplacements/Object 1"))
        frame.addElement(SvgTitle(text=replacement_title))
    return frame


# a writer table; a value is a string (one paragraph) or a list of elements for the cell
def table(rows, header_rows=0):
    t = OdfTable(name="Πίνακας1")
    t.addElement(TableColumn(numbercolumnsrepeated=len(rows[0])))
    header = TableHeaderRows()
    if header_rows:
        t.addElement(header)
    for i, row in enumerate(rows):
        tr = TableRow()
        for value in row:
            cell = TableCell(valuetype="string")
            if isinstance(value, str):
                cell.addElement(OdfP(text=value))
            else:
                inside(cell, *value)
            tr.addElement(cell)
        (header if i < header_rows else t).addElement(tr)
    return t


def bullets(*items):
    return inside(List(), *[inside(ListItem(), item if not isinstance(item, str) else OdfP(text=item)) for item in items])


def comment(text):
    return inside(Annotation(), Creator(text="Μαρία"), DcDate(text="2026-01-01T00:00:00"), OdfP(text=text))


def footnote(number, text):
    return inside(Note(id=f"ftn{number}", noteclass="footnote"), NoteCitation(text=number),
                  inside(NoteBody(), OdfP(text=text)))


def change_info():
    return inside(ChangeInfo(), Creator(text="Μαρία"), DcDate(text="2026-01-01T00:00:00"))


# one tracked change; libreoffice writes both ids
def region(change_id, change):
    return inside(ChangedRegion(qattributes={(XMLNS, "id"): change_id, (TEXTNS, "id"): change_id}), change)


# writes a text document with these elements in the body
def write_odt(path, *elements, styles=(), auto_styles=(), mimetype=None):
    doc = OpenDocumentText()
    if mimetype:
        doc.mimetype = mimetype
    doc.addPicture(PICTURE, "image/png", PNG)
    for style in styles:
        doc.styles.addElement(style)
    for style in auto_styles:
        doc.automaticstyles.addElement(style)
    for element in elements:
        doc.text.addElement(element)  # pyright: ignore[reportAttributeAccessIssue] (odfpy adds it at runtime)
    doc.save(str(path))
    return path


# the same document as make_docx in test_parsers.py, written by writer
def make_odt(path, mimetype=None):
    return write_odt(
        path,
        heading(1, "Εισαγωγή"),
        para("Πρώτη παράγραφος.", OdfLineBreak(), "Με δεύτερη γραμμή."),
        para(),
        para("   "),
        table([["Όνομα", "Βαθμός"], ["Μαρία", "9"]]),
        heading(2, "Μέθοδος"),
        para(image(desc="Αρχιτεκτονική του συστήματος")),
        para(image(desc="")),
        para("Εικόνα 1: Η ροή των δεδομένων", stylename="Caption"),
        heading(4, "Λεπτομέρειες"),
        para("Τελευταία παράγραφος."),
        mimetype=mimetype,
    )


# the same content gives the same text in odt and in docx
def test_read_odt_gives_the_same_text_as_docx(tmp_path):
    assert read_odf(make_odt(tmp_path / "doc.odt")) == EXPECTED_DOCX


NUMBERS = [["2023", "2024"], ["120", "340"]]

# (elements of the body, expected text); lambdas so every test gets new elements
ODT_CASES = {
    # headings
    "heading levels": (lambda: [heading(1, "Α"), heading(2, "Β"), heading(6, "Γ")], "# Α\n\n## Β\n\n###### Γ"),
    "empty heading is skipped": (lambda: [heading(1), para("Κείμενο")], "Κείμενο"),

    # text inside a paragraph
    "spaces and tabs": (lambda: [para("α", OdfS(c=3), "β", Tab(), "γ")], "α   β\tγ"),
    "text:s without a count is one space": (lambda: [para("α", OdfS(), "β")], "α β"),
    "spaces and new lines in the xml count as one space": (
        lambda: [para("Γεια\n        σου   κόσμε")], "Γεια σου κόσμε"),
    "spaces at the ends of a paragraph are removed": (lambda: [para("  Κείμενο  ")], "Κείμενο"),
    "non-breaking space is kept": (lambda: [para("Σύνολο: 10 €")], "Σύνολο: 10 €"),
    "links, spans and fields": (
        lambda: [para("Δες ", A(href="https://www.uoc.gr", type="simple", text="τον οδηγό"), ", το ",
                      Span(text="έντονο"), " κείμενο και την ", TextDate(text="11/10/2026"), ".")],
        "Δες τον οδηγό, το έντονο κείμενο και την 11/10/2026."),
    "soft page break inside a paragraph": (
        lambda: [para("Πρώτο μισό", SoftPageBreak(), " δεύτερο μισό")], "Πρώτο μισό δεύτερο μισό"),
    "soft page break between paragraphs": (lambda: [para("Α"), SoftPageBreak(), para("Β")], "Α\n\nΒ"),

    # lists and sections
    "list items are paragraphs": (lambda: [bullets("Πρώτο", "Δεύτερο")], "Πρώτο\n\nΔεύτερο"),
    "nested list keeps the order": (
        lambda: [inside(List(), inside(ListItem(), OdfP(text="α"), bullets("β")), inside(ListItem(), OdfP(text="γ")))],
        "α\n\nβ\n\nγ"),
    "sections and a section inside a section": (
        lambda: [inside(Section(name="Ενότητα1"), para("Α"), inside(Section(name="Ενότητα2"), para("Β"))),
                 para("Γ")],
        "Α\n\nΒ\n\nΓ"),

    # things that are not the text of the document
    "a comment is not text": (
        lambda: [para("Η προθεσμία είναι ", comment("να το ελέγξω"), "η Τρίτη.")], "Η προθεσμία είναι η Τρίτη."),
    "a comment on a range keeps the range": (
        lambda: [para("Η ", inside(Annotation(name="c1"), OdfP(text="σχόλιο")), "προθεσμία",
                      AnnotationEnd(name="c1"), " λήγει.")],
        "Η προθεσμία λήγει."),
    # same as docx and rtf: footnotes are left out and don't break the sentence
    "a footnote is not glued into the sentence": (
        lambda: [para("Κύρια πρόταση", footnote("1", "Πηγή: ΕΛΣΤΑΤ."), " συνεχίζει.")],
        "Κύρια πρόταση συνεχίζει."),
    "tracked changes give the current text": (
        lambda: [inside(TrackedChanges(),
                        region("ct1", inside(Deletion(), change_info(), OdfP(text="η Δευτέρα"))),
                        region("ct2", inside(Insertion(), change_info())),
                        region("ct3", inside(Deletion(), change_info(), OdfP(text="Διαγραμμένη παράγραφος.")))),
                 para("Η προθεσμία είναι ", Change(changeid="ct1"), ChangeStart(changeid="ct2"), "η Τρίτη",
                      ChangeEnd(changeid="ct2"), "."),
                 para(Change(changeid="ct3"))],
        "Η προθεσμία είναι η Τρίτη."),
    # the table of contents repeats the headings
    "table of contents is skipped": (
        lambda: [inside(TableOfContent(name="Περιεχόμενα1"),
                        inside(IndexBody(), inside(IndexTitle(name="Περιεχόμενα1_Head"), OdfP(text="Περιεχόμενα")),
                               OdfP(text="Εισαγωγή\t1"))),
                 heading(1, "Εισαγωγή")],
        "# Εισαγωγή"),

    # images, text boxes and shapes
    "image alt text from the title": (lambda: [para(image(title="Χάρτης"))], "[Image: Χάρτης]"),
    "image alt text from the description": (lambda: [para(image(desc="Χάρτης"))], "[Image: Χάρτης]"),
    "image without alt text": (lambda: [para(image()), para("Κείμενο")], "Κείμενο"),
    "image with a title and a description uses the description": (
        lambda: [para(image(title="Χάρτης", desc="Χάρτης της Κρήτης"))], "[Image: Χάρτης της Κρήτης]"),
    "a text box with a name is not an image": (
        lambda: [para(inside(textbox(para("Σημείωση")), SvgTitle(text="Πλαίσιο1")))], "Σημείωση"),
    "image with only spaces as alt text": (lambda: [para(image(title="  ", desc="  ")), para("Κείμενο")],
                                           "Κείμενο"),
    # like docx, the image comes before the text of its paragraph
    "image inside a sentence": (
        lambda: [para("Δες την εικόνα", image(desc="Χάρτης"), " παρακάτω.")],
        "[Image: Χάρτης]\n\nΔες την εικόνα παρακάτω."),
    "image with a caption": (
        lambda: [para(textbox(para(image(desc="Χάρτης"), "Εικόνα ",
                                   Sequence(name="Illustration", formula="ooow:Illustration+1", numformat="1",
                                            text="1"),
                                   ": Χάρτης Κρήτης")))],
        "[Image: Χάρτης]\n\nΕικόνα 1: Χάρτης Κρήτης"),
    "text box appears once, in its place": (
        lambda: [para("Πριν"), para(textbox(para("Κείμενο σε πλαίσιο"))), para("Μετά")],
        "Πριν\n\nΚείμενο σε πλαίσιο\n\nΜετά"),
    "text box with two paragraphs": (lambda: [para(textbox(para("α"), para("β")))], "α\n\nβ"),
    "empty text box": (lambda: [para(textbox(para())), para("Κείμενο")], "Κείμενο"),
    "text box anchored to the page": (
        lambda: [textbox(para("Σημείωση"), anchortype="page"), para("Κείμενο")], "Σημείωση\n\nΚείμενο"),
    "shape with text": (
        lambda: [para(inside(CustomShape(width="3cm", height="1cm", anchortype="paragraph"), para("Βέλος")))],
        "Βέλος"),
    "group of shapes": (
        lambda: [para(inside(G(anchortype="paragraph"),
                             inside(CustomShape(width="3cm", height="1cm"), para("α")),
                             inside(CustomShape(width="3cm", height="1cm"), para("β"))))],
        "α\n\nβ"),

    # tables
    "table between paragraphs": (
        lambda: [para("Πριν"), table([["Όνομα", "Βαθμός"], ["Μαρία", "9"]]), para("Μετά")],
        "Πριν\n\nΌνομα: Μαρία; Βαθμός: 9\n\nΜετά"),
    "header row marked by the file": (lambda: [table(NUMBERS, header_rows=1)], "2023: 120; 2024: 340"),
    "no header row marked": (lambda: [table(NUMBERS)], "2023 | 2024\n120 | 340"),
    "cell with two paragraphs": (
        lambda: [table([["Όνομα", "Σημειώσεις"], ["Μαρία", [OdfP(text="α"), OdfP(text="β")]]])],
        "Όνομα: Μαρία; Σημειώσεις: α / β"),
    "cell with a list": (
        lambda: [table([["Όνομα", "Σημειώσεις"], ["Μαρία", [bullets("α", "β")]]])],
        "Όνομα: Μαρία; Σημειώσεις: α / β"),
    "empty table": (lambda: [table([["", ""], ["", ""]]), para("Κείμενο")], "Κείμενο"),
    "image in a table cell": (
        lambda: [table([["Όνομα", "Φωτογραφία"], ["Μαρία", [para(image(desc="Πορτρέτο"))]]])],
        "Όνομα: Μαρία; Φωτογραφία: [Image: Πορτρέτο]"),
    # a cell value is not a heading, the # marks belong only to blocks
    "heading in a table cell has no #": (
        lambda: [table([["Όνομα", "Ενότητα"], ["Μαρία", [heading(2, "Τίτλος")]]])],
        "Όνομα: Μαρία; Ενότητα: Τίτλος"),

    # more structure
    "heading level 0 is level 1": (lambda: [heading(0, "Μηδέν")], "# Μηδέν"),
    "list header is read like a list item": (
        lambda: [inside(List(), inside(ListHeader(), OdfP(text="Κεφαλίδα λίστας")),
                        inside(ListItem(), OdfP(text="Πρώτο")))],
        "Κεφαλίδα λίστας\n\nΠρώτο"),
    "numbered paragraph outside a list": (
        lambda: [inside(NumberedParagraph(listid="L1"), OdfP(text="Αριθμημένη"))], "Αριθμημένη"),
    "alphabetical index is skipped": (
        lambda: [inside(AlphabeticalIndex(name="Ευρετήριο1"), inside(IndexBody(), OdfP(text="λέξη, 3"))),
                 para("Κείμενο")],
        "Κείμενο"),
    # unlike the other indexes, the bibliography is the only place the full references are written
    "bibliography is kept": (
        lambda: [inside(Bibliography(name="Βιβλιογραφία1"), inside(IndexBody(), OdfP(text="[1] Παπαδάκης, 2020"))),
                 para("Κείμενο")],
        "[1] Παπαδάκης, 2020\n\nΚείμενο"),

    # objects and links around images
    "chart is skipped": (lambda: [para("Πριν"), para(chart()), para("Μετά")], "Πριν\n\nΜετά"),
    "chart alt text is kept like an image": (
        lambda: [para(chart(replacement_title="Εγγραφές ανά έτος"))], "[Image: Εγγραφές ανά έτος]"),
    "image inside a link": (
        lambda: [para("Δες ", inside(DrawA(href="https://www.uoc.gr", type="simple"), image(desc="Λογότυπο")))],
        "[Image: Λογότυπο]\n\nΔες"),
}


# odt files as writer saves them
@pytest.mark.parametrize("elements, expected", ODT_CASES.values(), ids=ODT_CASES.keys())
def test_read_odt(tmp_path, elements, expected):
    assert read_odf(write_odt(tmp_path / "doc.odt", *elements())) == expected


# a table inside a table cell (common in forms)
def test_read_odt_table_inside_a_table_cell(tmp_path):
    inner = table([["Τηλέφωνο", "Email"], ["2810123456", "maria@uoc.gr"]])
    f = write_odt(tmp_path / "nested.odt", table([["Στοιχεία", [inner]]]))

    text = read_odf(f)

    assert "2810123456" in text
    assert "maria@uoc.gr" in text


# the Title style is a level 1 heading like in docx, also through the styles that are based on it
@pytest.mark.parametrize("style, styles, auto_styles, expected", [
    pytest.param("Title", [Style(name="Title", family="paragraph")], [], "# Τίτλος", id="Title"),
    pytest.param("P1", [Style(name="Title", family="paragraph")],
                 [Style(name="P1", family="paragraph", parentstylename="Title")], "# Τίτλος",
                 id="automatic style based on Title"),
    pytest.param("P1", [Style(name="Title", family="paragraph"),
                        Style(name="Κεφάλαιο", family="paragraph", parentstylename="Title")],
                 [Style(name="P1", family="paragraph", parentstylename="Κεφάλαιο")], "# Τίτλος",
                 id="based on a style based on Title"),
    pytest.param("Subtitle", [Style(name="Subtitle", family="paragraph")], [], "Τίτλος", id="Subtitle is text"),
])
def test_read_odt_title_style(tmp_path, style, styles, auto_styles, expected):
    f = write_odt(tmp_path / "doc.odt", para("Τίτλος", stylename=style), para("Κείμενο"),
                  styles=styles, auto_styles=auto_styles)

    assert read_odf(f) == expected + "\n\nΚείμενο"


# a heading style can say its level with default-outline-level, with its name (Heading_20_N), or both
HEADING_STYLES = {
    "level": lambda n: Style(name=f"Κεφάλαιο_20_{n}", displayname=f"Κεφάλαιο {n}", family="paragraph",
                             defaultoutlinelevel=str(n)),
    "name": lambda n: Style(name=f"Heading_20_{n}", displayname=f"Heading {n}", family="paragraph"),
    "level and name": lambda n: Style(name=f"Heading_20_{n}", displayname=f"Heading {n}", family="paragraph",
                                      defaultoutlinelevel=str(n)),
}


# google docs writes headings as text:p with an automatic style based on the heading style, not as text:h
@pytest.mark.parametrize("kind, automatic, level", [
    pytest.param(kind, automatic, level, id=f"{kind}, {'automatic style' if automatic else 'the style itself'}, {level}")
    for kind, automatic, level in itertools.product(HEADING_STYLES, [False, True], [1, 3])
])
def test_read_odt_heading_style(tmp_path, kind, automatic, level):
    style = HEADING_STYLES[kind](level)
    name = style.getAttribute("name")
    auto_styles = [Style(name="P64", family="paragraph", parentstylename=name)] if automatic else []
    f = write_odt(tmp_path / "doc.odt", para("Μέθοδος", stylename="P64" if automatic else name), para("Κείμενο"),
                  styles=[style], auto_styles=auto_styles)

    assert read_odf(f) == "#" * level + " Μέθοδος\n\nΚείμενο"


# which level wins when a style says more than one thing
@pytest.mark.parametrize("element, styles, expected", [
    pytest.param(lambda: para("Μέθοδος", stylename="Heading"), lambda: [Style(name="Heading", family="paragraph")],
                 "Μέθοδος", id="Heading without a number is text"),
    pytest.param(lambda: para("Μέθοδος", stylename="Heading_20_2_20_old"),
                 lambda: [Style(name="Heading_20_2_20_old", family="paragraph")], "Μέθοδος",
                 id="a name only partly like a heading is text"),
    pytest.param(lambda: para("Μέθοδος", stylename="Heading_20_2"),
                 lambda: [Style(name="Heading_20_2", family="paragraph", defaultoutlinelevel="3")], "### Μέθοδος",
                 id="the written level wins over the name"),
    pytest.param(lambda: para("Μέθοδος", stylename="Κεφάλαιο"),
                 lambda: [HEADING_STYLES["level and name"](2),
                          Style(name="Κεφάλαιο", family="paragraph", parentstylename="Heading_20_2",
                                defaultoutlinelevel="3")],
                 "### Μέθοδος", id="the level of a style wins over its parent's"),
    pytest.param(lambda: para("Μέθοδος", cls=H, outlinelevel=3, stylename="Heading_20_2"),
                 lambda: [HEADING_STYLES["level and name"](2)], "### Μέθοδος",
                 id="the level of text:h wins over its style"),
    pytest.param(lambda: table([["Όνομα", "Ενότητα"], ["Μαρία", [para("Μέθοδος", stylename="Heading_20_2")]]]),
                 lambda: [HEADING_STYLES["level and name"](2)], "Όνομα: Μαρία; Ενότητα: Μέθοδος",
                 id="heading style in a table cell has no #"),
])
def test_read_odt_heading_style_priority(tmp_path, element, styles, expected):
    assert read_odf(write_odt(tmp_path / "doc.odt", element(), styles=styles())) == expected


# a name is unique only inside its family: a text style based on Title with the name of a paragraph style
# (odfpy renames such a style when it saves, so the name is changed in the xml afterwards)
def test_read_odt_text_style_does_not_make_a_title(tmp_path):
    plain = write_odt(tmp_path / "plain.odt", para(Span(stylename="T", text="Μέθοδος"), stylename="P64"),
                      styles=[Style(name="Title", family="paragraph"), Style(name="Κεφάλαιο", family="paragraph")],
                      auto_styles=[Style(name="P64", family="paragraph", parentstylename="Κεφάλαιο"),
                                   Style(name="T", family="text", parentstylename="Title")])
    f = tmp_path / "doc.odt"
    with zipfile.ZipFile(plain) as zin, zipfile.ZipFile(f, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "content.xml":
                assert b'style:name="T"' in data
                data = data.replace(b'"T"', '"Κεφάλαιο"'.encode())
            zout.writestr(item, data)

    assert read_odf(f) == "Μέθοδος"


# page header, footer and document properties are not body text
def test_read_odt_header_footer_and_properties_are_not_text(tmp_path):
    doc = OpenDocumentText()
    layout = PageLayout(name="pm1")
    doc.automaticstyles.addElement(layout)
    page = MasterPage(name="Standard", pagelayoutname=layout)
    page.addElement(inside(Header(), OdfP(text="Πανεπιστήμιο Κρήτης")))
    page.addElement(inside(Footer(), OdfP(text="Σελίδα 1")))
    doc.masterstyles.addElement(page)
    doc.meta.addElement(DcTitle(text="Μυστικός τίτλος"))
    doc.text.addElement(OdfP(text="Κείμενο"))  # pyright: ignore[reportAttributeAccessIssue] (odfpy adds it at runtime)
    f = tmp_path / "doc.odt"
    doc.save(str(f))

    assert read_odf(f) == "Κείμενο"


# some programs put spaces and new lines between the tags of content.xml
def test_read_odt_pretty_printed_xml(tmp_path):
    plain = write_odt(tmp_path / "plain.odt", heading(1, "Εισαγωγή"), para("Κείμενο."),
                      table([["Όνομα", "Βαθμός"], ["Μαρία", "9"]]))
    pretty = tmp_path / "pretty.odt"
    with zipfile.ZipFile(plain) as zin, zipfile.ZipFile(pretty, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "content.xml":
                data = re.sub(rb">\s*<(?!/text:[ph]>)", b">\n   <", data)
            zout.writestr(item, data)

    assert read_odf(pretty) == read_odf(plain) == "# Εισαγωγή\n\nΚείμενο.\n\nΌνομα: Μαρία; Βαθμός: 9"


# a document with nothing in it gives empty text
def test_read_odt_empty_document(tmp_path):
    assert read_odf(write_odt(tmp_path / "empty.odt")) == ""


# a spreadsheet renamed to .odt is not a text document
def test_read_odf_refuses_a_spreadsheet(tmp_path):
    doc = OpenDocumentSpreadsheet()
    doc.spreadsheet.addElement(OdfTable(name="Φύλλο"))  # pyright: ignore[reportAttributeAccessIssue] (odfpy adds it at runtime)
    f = tmp_path / "sheet.odt"
    doc.save(str(f))

    with pytest.raises(ValueError):
        read_odf(f)


# odt and its template point to read_odf
@pytest.mark.parametrize("ext", [".odt", ".ott"])
def test_odt_extensions_are_registered(ext):
    assert PARSERS.get(ext) is read_odf


# odt, its template and an upper case name through parse_file
@pytest.mark.parametrize("name, mimetype", [
    ("doc.odt", None),
    ("template.ott", "application/vnd.oasis.opendocument.text-template"),
    ("REPORT.ODT", None),
])
def test_parse_file_reads_odt(tmp_path, name, mimetype):
    doc = parse_one(tmp_path, name, built(tmp_path, "base.odt", lambda f: make_odt(f, mimetype)))

    assert doc.parse_status == "ok"
    assert doc.content_type == "prose"
    assert doc.raw_text == EXPECTED_DOCX


# a broken odt is failed, not a crash
def docx_bytes(tmp_path):
    d = docx.Document()
    d.add_paragraph("Κείμενο")
    return built(tmp_path, "base.docx", lambda f: d.save(str(f)))


def zip_without_content(tmp_path):
    f = tmp_path / "base.zip"
    with zipfile.ZipFile(f, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
    return f.read_bytes()


@pytest.mark.parametrize("make", [
    pytest.param(lambda tmp_path: b"this is not a zip file", id="not a zip"),
    pytest.param(docx_bytes, id="a docx renamed to odt"),
    pytest.param(zip_without_content, id="zip without content.xml"),
])
def test_parse_file_marks_a_broken_odt_as_failed(tmp_path, make):
    doc = parse_one(tmp_path, "broken.odt", make(tmp_path))

    assert doc.parse_status == "failed"
    assert doc.content_type == "prose"


# an odt with no text is empty, not failed
def test_parse_file_odt_without_text_is_empty(tmp_path):
    data = built(tmp_path, "base.odt", lambda f: write_odt(f, para(), para("   "), para(image())))

    assert parse_one(tmp_path, "doc.odt", data).parse_status == "empty"


# ---------------- .odp ----------------

# a frame with a text box on a slide; cls is the presentation class (title, outline, subtitle, notes)
def slide_box(*paragraphs, cls=None):
    frame = Frame(**{"class": cls}) if cls else Frame()
    return inside(frame, inside(TextBox(), *[OdfP(text=p) if isinstance(p, str) else p for p in paragraphs]))


def speaker_notes(*paragraphs):
    return inside(Notes(), PageThumbnail(pagenumber=1), slide_box(*paragraphs, cls="notes"))


# writes a presentation; every page is a list of shapes, hidden holds the numbers of the hidden pages
def write_odp(path, *pages, hidden=(), master_text=None, footer=None, mimetype=None):
    doc = OpenDocumentPresentation()
    if mimetype:
        doc.mimetype = mimetype
    doc.addPicture(PICTURE, "image/png", PNG)
    layout = PageLayout(name="PL")
    doc.automaticstyles.addElement(layout)
    master = MasterPage(name="Default", pagelayoutname=layout)
    if master_text:
        master.addElement(slide_box(master_text, cls="title"))
    doc.masterstyles.addElement(master)
    hidden_style = inside(Style(name="dp1", family="drawing-page"), DrawingPageProperties(visibility="hidden"))
    doc.automaticstyles.addElement(hidden_style)
    if footer:
        doc.presentation.addElement(FooterDecl(name="ftr1", text=footer))  # pyright: ignore[reportAttributeAccessIssue] (odfpy adds it at runtime)
    for number, shapes in enumerate(pages, start=1):
        page = Page(masterpagename=master, stylename=hidden_style) if number in hidden else Page(masterpagename=master)
        doc.presentation.addElement(inside(page, *shapes))  # pyright: ignore[reportAttributeAccessIssue] (odfpy adds it at runtime)
    doc.save(str(path))
    return path


# the same presentation as make_pptx in test_parsers.py, written by impress
def make_odp(path, mimetype=None):
    return write_odp(
        path,
        [slide_box("  Εισαγωγή ", cls="title"), slide_box(bullets("Πρώτη κουκκίδα", "Δεύτερη κουκκίδα"), cls="outline"),
         slide_box(), speaker_notes("να πω για το dataset")],
        [inside(Frame(), table([["Όνομα", "Βαθμός"], ["Μαρία", "9"]])), image(desc="Διάγραμμα"), image(desc=""),
         inside(G(), slide_box("Κείμενο μέσα σε ομάδα"), inside(G(), slide_box("Ομάδα μέσα σε ομάδα")))],
        [slide_box(cls="title"), speaker_notes("   ")],
        mimetype=mimetype,
    )


# the same content gives the same text in odp and in pptx
def test_read_odp_gives_the_same_text_as_pptx(tmp_path):
    assert read_odf(make_odp(tmp_path / "talk.odp")) == EXPECTED_PPTX


# one slide: (shapes, the text after "# Slide 1")
ODP_CASES = {
    "no title": ([slide_box("κείμενο")], "\n\nκείμενο"),
    "title with a line break": ([slide_box(para("Γραμμή 1", OdfLineBreak(), "Γραμμή 2"), cls="title")],
                                ": Γραμμή 1 Γραμμή 2"),
    "title in two paragraphs": ([slide_box("Γραμμή 1", "  ", "Γραμμή 2", cls="title")], ": Γραμμή 1 Γραμμή 2"),
    "title after other shapes": ([slide_box("κείμενο"), slide_box("Τίτλος", cls="title")], ": Τίτλος\n\nκείμενο"),
    "only the first title is the title": ([slide_box("Α", cls="title"), slide_box("Β", cls="title")], ": Α\n\nΒ"),
    "an empty title is still the title": ([slide_box(cls="title"), slide_box("Β", cls="title")], "\n\nΒ"),
    "subtitle is slide text": ([slide_box("Τίτλος", cls="title"), slide_box("υπότιτλος", cls="subtitle")],
                               ": Τίτλος\n\nυπότιτλος"),
    "paragraphs and a line break in one box": ([slide_box("α", para("β", OdfLineBreak(), "γ"))], "\n\nα\nβ\nγ"),
    "empty paragraphs in a box": ([slide_box("α", "", "   ", "β")], "\n\nα\nβ"),
    "list inside a list": ([slide_box(bullets("α", bullets("β", "γ")))], "\n\nα\nβ\nγ"),
    "a heading in a box is not a heading": ([slide_box(heading(1, "Ενότητα"))], "\n\nΕνότητα"),
    "text directly in a shape": ([inside(CustomShape(), OdfP(text="κείμενο σε σχήμα"))], "\n\nκείμενο σε σχήμα"),
    "two shapes are two blocks": ([slide_box("α"), slide_box("β")], "\n\nα\n\nβ"),
    "image with a title only": ([image(title="Λογότυπο")], "\n\n[Image: Λογότυπο]"),
    "image without alt text": ([image()], ""),
    "chart": ([chart()], ""),
    "notes in two paragraphs and a line break": (
        [slide_box("Τίτλος", cls="title"), speaker_notes("α", para("β", OdfLineBreak(), "γ"))],
        ": Τίτλος\n\nNotes: α\nβ\nγ"),
    "notes only with spaces": ([slide_box("Τίτλος", cls="title"), speaker_notes("  ", "")], ": Τίτλος"),
}


# every kind of shape on one slide
@pytest.mark.parametrize("shapes, expected", ODP_CASES.values(), ids=ODP_CASES.keys())
def test_read_odp(tmp_path, shapes, expected):
    assert read_odf(write_odp(tmp_path / "talk.odp", shapes)) == "# Slide 1" + expected


# a comment on a slide is not text; older libreoffice writes it with another prefix (odfpy can only write the new one)
@pytest.mark.parametrize("comment_xml", [
    "<office:annotation><dc:creator>M</dc:creator><text:p>σχόλιο</text:p></office:annotation>",
    '<officeooo:annotation xmlns:officeooo="http://openoffice.org/2009/office"><dc:creator>M</dc:creator>'
    "<text:p>σχόλιο</text:p></officeooo:annotation>",
], ids=["office", "officeooo"])
def test_read_odp_slide_comment_is_not_text(tmp_path, comment_xml):
    plain = write_odp(tmp_path / "plain.odp", [slide_box("κείμενο")])
    f = tmp_path / "comment.odp"
    with zipfile.ZipFile(plain) as zin, zipfile.ZipFile(f, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "content.xml":
                data = data.replace(b"</draw:page>", comment_xml.encode() + b"</draw:page>")
            zout.writestr(item, data)

    assert read_odf(f) == "# Slide 1\n\nκείμενο"


# slides are numbered in order, empty and hidden ones too
def test_read_odp_counts_every_slide(tmp_path):
    f = write_odp(tmp_path / "talk.odp", [slide_box("Α", cls="title")], [], [slide_box("Γ", cls="title")], hidden=(3,))

    assert read_odf(f) == "# Slide 1: Α\n\n# Slide 2\n\n# Slide 3: Γ"


# the master page, the page picture of the notes and the footer are not slide text
def test_read_odp_master_and_footer_are_not_text(tmp_path):
    f = write_odp(tmp_path / "talk.odp", [slide_box("κείμενο")], master_text="Όνομα μαθήματος", footer="Σελίδα")

    assert read_odf(f) == "# Slide 1\n\nκείμενο"


# some programs put spaces and new lines between the tags of content.xml
def test_read_odp_pretty_printed_xml(tmp_path):
    plain = make_odp(tmp_path / "plain.odp")
    pretty = tmp_path / "pretty.odp"
    with zipfile.ZipFile(plain) as zin, zipfile.ZipFile(pretty, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "content.xml":
                data = re.sub(rb">\s*<(?!/text:[ph]>)", b">\n   <", data)
            zout.writestr(item, data)

    assert read_odf(pretty) == EXPECTED_PPTX


# a presentation without slides gives empty text
def test_read_odp_without_slides(tmp_path):
    assert read_odf(write_odp(tmp_path / "empty.odp")) == ""


# odp and its template point to read_odf
@pytest.mark.parametrize("ext", [".odp", ".otp"])
def test_odp_extensions_are_registered(ext):
    assert PARSERS.get(ext) is read_odf


# odp, its template, an upper case name, and an odp saved as .odt through parse_file
@pytest.mark.parametrize("name, mimetype", [
    ("talk.odp", None),
    ("template.otp", "application/vnd.oasis.opendocument.presentation-template"),
    ("TALK.ODP", None),
    ("talk.odt", None),
])
def test_parse_file_reads_odp(tmp_path, name, mimetype):
    doc = parse_one(tmp_path, name, built(tmp_path, "base.odp", lambda f: make_odp(f, mimetype)))

    assert doc.parse_status == "ok"
    assert doc.content_type == "prose"
    assert doc.raw_text == EXPECTED_PPTX


# a broken odp is failed, not a crash
def test_parse_file_marks_a_broken_odp_as_failed(tmp_path):
    doc = parse_one(tmp_path, "broken.odp", b"this is not a zip file")

    assert doc.parse_status == "failed"
    assert doc.content_type == "prose"
