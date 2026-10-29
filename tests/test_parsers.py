import datetime
import itertools
import logging
import os
import re
import zipfile
from pathlib import Path

import docx
import openpyxl
from openpyxl.worksheet.hyperlink import Hyperlink
import pandas as pd
import pptx
import pytest
from docx.opc.constants import RELATIONSHIP_TYPE as DOCX_RT
from docx.oxml import parse_xml
from docx.shared import Inches
from odf.draw import Frame, Page, TextBox as OdfTextBox
from odf.opendocument import OpenDocumentPresentation, OpenDocumentSpreadsheet, OpenDocumentText
from odf.table import CoveredTableCell, Table as OdfTable, TableCell, TableRow
from odf.text import A as OdfA, LineBreak as OdfLineBreak, P as OdfP, Span as OdfSpan
from PIL import Image
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.shapes.autoshape import Shape as PptxShape
from pptx.opc.constants import RELATIONSHIP_TYPE as PPTX_RT
from pptx.oxml.ns import qn as pptx_qn
from pptx.util import Inches as PptxInches

from rag_fs.config import CorpusConfig
from rag_fs.ingest.file_scanner import scan
from rag_fs.ingest.parsers import PARSERS, parse_file, tabular
from rag_fs.ingest.parsers.office import (
    format_shape, heading_level, read_docx, read_pptx, read_rtf, read_slide_variants,
    read_word_variants,
)
from rag_fs.ingest.parsers.opendocument import read_odf, read_ods
from rag_fs.ingest.parsers.tabular import cell_to_str, fill_merged, format_rows, read_csv, read_xls, read_xlsx
from rag_fs.ingest.parsers.text import read_text

FIXTURES = Path(__file__).parent / "fixtures"

ALLOWED_STATUSES = {"ok", "ocr", "empty", "failed", "unsupported"}
ALLOWED_CONTENT_TYPES = {"prose", "code", "tabular", "binary"}


# writes the files into a corpus folder and parses them
def parse_all(tmp_path, files):
    root = tmp_path / "corpus"
    root.mkdir()
    for name, data in files.items():
        (root / name).write_bytes(data)
    result = scan(CorpusConfig(roots=[root]))
    docs = {}
    for scanned in result.files:
        docs[scanned.rel_path.name] = parse_file(scanned)
    return docs


# same as parse_all for a single file
def parse_one(tmp_path, name, data):
    return parse_all(tmp_path, {name: data})[name]


# makes a test file outside the corpus and returns its bytes
def built(tmp_path, name, make):
    f = tmp_path / name
    make(f)
    return f.read_bytes()


# a small corpus with one file for each case parse_file has to handle
@pytest.fixture
def docs(tmp_path):
    return parse_all(tmp_path, {
        "notes.txt": "Γεια\r\nσου\rκόσμε".encode("utf-8"),
        "main.py": b"print(1)\n",
        "blank.md": b"  \n\n",
        "empty.txt": b"",
        "data.zip": b"PK\x03\x04",
        "fake.txt": b"ELF\x00\x00\x01",
        "Makefile": b"all:\n\tgcc a.c\n",
        "weird.xyz": bytes(range(256)) * 4,
        "old.txt": "Λογαριασμός ρεύματος".encode("cp1253"),
    })


# status, content type and text for each file (None = not checked)
@pytest.mark.parametrize("name, status, content_type, text", [
    ("notes.txt", "ok", "prose", "Γεια\nσου\nκόσμε"),        # line endings normalized
    ("main.py", "ok", "code", None),
    ("blank.md", "empty", None, None),                       # whitespace only
    ("empty.txt", "empty", None, None),                      # 0 bytes
    ("data.zip", "unsupported", "binary", ""),               # known binary extension
    ("fake.txt", "failed", "prose", ""),                     # known text extension, binary inside
    ("weird.xyz", "unsupported", "binary", None),            # unknown extension, binary inside
    ("Makefile", "ok", None, "all:\n\tgcc a.c\n"),           # no extension, read as text
    ("old.txt", None, None, "Λογαριασμός ρεύματος"),         # old Greek encoding
])
def test_parse_file_status_type_and_text(docs, name, status, content_type, text):
    doc = docs[name]
    if status is not None:
        assert doc.parse_status == status
    if content_type is not None:
        assert doc.content_type == content_type
    if text is not None:
        assert doc.raw_text == text


# more code extensions: c++ headers, shaders, yacc/lex, assembly, build and shell scripts, other languages
@pytest.mark.parametrize("name", [
    "vec.hpp", "vec.hh", "vec.hxx", "vec.cc", "vec.cxx", "vec.inl",
    "light.hlsl", "light.glsl", "light.shader", "light.cginc", "light.compute", "light.frag", "light.vert",
    "parser.y", "lexer.l", "boot.asm", "start.s", "tools.cmake",
    "run.bat", "run.cmd", "run.ps1", "init.lua", "tool.pl", "Main.scala", "build.kts", "view.mm",
])
def test_more_code_extensions(tmp_path, name):
    doc = parse_one(tmp_path, name, b"int x = 1;\n")
    assert doc.content_type == "code"
    assert doc.raw_text == "int x = 1;\n"


# a file with no extension (Makefile) gets ext ""
def test_file_without_extension_has_an_empty_ext(docs):
    assert docs["Makefile"].ext == ""


# parse_file never makes up a status or a content type
def test_only_known_labels_are_used(docs):
    for doc in docs.values():
        assert doc.parse_status in ALLOWED_STATUSES
        assert doc.content_type in ALLOWED_CONTENT_TYPES


# path, root and hash in the Document are the ones the scanner found
def test_document_fields_come_from_the_scanned_file(tmp_path):
    root = tmp_path / "corpus"
    (root / "sub").mkdir(parents=True)
    (root / "sub" / "a.txt").write_text("hello")
    scanned = scan(CorpusConfig(roots=[root])).files[0]

    doc = parse_file(scanned)

    assert doc.file_path == Path("sub/a.txt")
    assert doc.root_id == scanned.root_id
    assert doc.sha256 == scanned.sha256
    assert doc.ext == ".txt"


# a file that can't be parsed is logged and marked failed
def test_failure_is_logged_not_raised(tmp_path, caplog):
    with caplog.at_level(logging.WARNING):
        doc = parse_one(tmp_path, "fake.txt", b"\x00\x01\x02")

    assert doc.parse_status == "failed"
    assert "fake.txt" in caplog.text


# the date of the file is kept even when it fails or is unsupported
@pytest.mark.parametrize("name, data", [
    ("notes.txt", "Γεια".encode()),        # read fine
    ("broken.docx", b"not a zip"),         # failed: the date must still be there
    ("song.mp3", b"\x00\x01"),             # unsupported: the date must still be there
])
def test_parse_file_keeps_the_last_modified_time(tmp_path, name, data):
    root = tmp_path / "corpus"
    root.mkdir()
    f = root / name
    f.write_bytes(data)
    when = datetime.datetime(2024, 3, 1, 14, 30).timestamp()
    os.utime(f, (when, when))   # (access time, modification time)

    doc = parse_file(scan(CorpusConfig(roots=[root])).files[0])

    assert doc.last_modified == when


# REPORT.DOCX is read like report.docx
def test_parse_file_reads_upper_case_extensions(tmp_path):
    d = docx.Document()
    d.add_paragraph("Αναφορά")
    data = built(tmp_path, "base.docx", lambda f: d.save(str(f)))

    doc = parse_one(tmp_path, "REPORT.DOCX", data)

    assert doc.parse_status == "ok"
    assert doc.raw_text == "Αναφορά"


# one read_text case: the text written in the given encoding
# mostly english with one greek sentence: the english must not make it look like broken utf-8
ENGLISH_THEN_GREEK = ("Please send the report about the exams to the secretary before Friday so that everyone has time "
                      "to read it carefully and prepare questions for the meeting next week. ") * 3 + \
    "Θέμα: η εξεταστική του Ιουνίου στο αμφιθέατρο Α."


def encoded(text, encoding, marks=()):
    return pytest.param(text.encode(encoding), text, marks=marks, id=f"{encoding}: {text[:25]}")


# greek text with english terms and the opposite
MIXED_GREEK_ENGLISH = [
    "The retrieval step uses BM25 and dense embeddings. "
    "In the thesis we call it ανάκτηση πληροφορίας (information retrieval).",
    "Η ανάκτηση γίνεται με BM25 και dense embeddings, και μετά re-ranking με cross-encoder.",
]


# read_text finds the right encoding
@pytest.mark.parametrize("data, expected", [
    encoded("Λογαριασμός ρεύματος Μαρτίου: 54€", "utf-8"),
    encoded("Ο λογαριασμός του ρεύματος για τον Μάρτιο ήταν 54 ευρώ.", "cp1253"),
    encoded("Λογαριασμός ρεύματος", "cp1253"),
    encoded("Σημειώσεις", "cp1253"),
    encoded("Café crème, déjà vu à Paris.", "cp1252"),
    encoded("Γεια σου κόσμε", "utf-16"),
    encoded("Ναι", "cp1253", marks=pytest.mark.xfail(strict=True, reason="L28: too short to detect cp1253")),
    encoded("Όχι", "cp1253"),
    encoded("Άρτα", "cp1253"),
    encoded("Χανιά", "cp1253"),
    encoded("Ηράκλειο", "cp1253"),
    encoded("Άνοιξη στην Ήπειρο, Ώρα για Ύδρα.", "iso8859_7"),
    *[encoded(text, enc) for text in MIXED_GREEK_ENGLISH for enc in ("utf-8", "cp1253", "iso8859_7")],
    encoded("Πήγα στην Άρτα και μετά στην Αθήνα με τα παιδιά μου.", "iso8859_7"),
    encoded("Παράγραφος ¶ 3 του κανονισμού σπουδών.", "cp1253"),
    pytest.param("¶Ενότητα 3".encode(), "¶Ενότητα 3", id="utf-8 with a pilcrow before greek"),
    pytest.param("αβγδεζηθικ".encode() + b"\x93", "αβγδεζηθικ\ufffd", id="utf-8 with one bad byte in ten greek letters"),
    pytest.param(b"Hello \x93world\x94", "Hello \u201cworld\u201d", id="english with windows quotes"),
    pytest.param(ENGLISH_THEN_GREEK.encode("cp1253"), ENGLISH_THEN_GREEK, id="cp1253: long english, one greek sentence"),
    pytest.param("\ufeffαβγδεζηθικ".encode() + b"\x93", "αβγδεζηθικ\ufffd", id="BOM and one bad byte"),
    pytest.param("\ufeffΓεια σου".encode("utf-8"), "Γεια σου", id="BOM removed"),
    pytest.param(b"", "", id="empty file"),
    pytest.param(b"a\r\nb", "a\r\nb", id="line endings kept (parse_file normalizes them)"),
])
def test_read_text(tmp_path, data, expected):
    f = tmp_path / "file.txt"
    f.write_bytes(data)

    assert read_text(f) == expected


# a NUL byte means binary, also when the file is utf-8 with a few broken bytes
def test_read_text_nul_byte_in_utf8_with_a_broken_byte(tmp_path):
    f = tmp_path / "data.bin"
    f.write_bytes("Γεια σου κόσμε ".encode() * 5 + b"\x00\x93")

    with pytest.raises(ValueError):
        read_text(f)


# a few windows bytes in a utf-8 file should not lose the greek
def test_read_text_utf8_with_stray_windows_quotes_keeps_the_greek(tmp_path):
    f = tmp_path / "mixed.txt"
    f.write_bytes("Σημειώσεις για την εξεταστική ".encode() + b"\x93quote\x94"
                  + " και τα θέματα του Ιουνίου.".encode())

    text = read_text(f)

    assert "Σημειώσεις για την εξεταστική" in text
    assert "θέματα του Ιουνίου" in text


# binary data is not accepted as text
@pytest.mark.parametrize("data", [bytes(range(256)) * 4, b"ELF\x00\x00\x01abc"])
def test_read_text_rejects_binary(tmp_path, data):
    f = tmp_path / "file.bin"
    f.write_bytes(data)

    with pytest.raises(ValueError):
        read_text(f)


# how one cell value becomes text
@pytest.mark.parametrize("value, expected", [
    (None, ""),
    (datetime.datetime(2024, 3, 1), "2024-03-01"),
    (datetime.datetime(2024, 3, 1, 14, 30), "2024-03-01 14:30:00"),
    (54, "54"),
    (54.0, "54"),
    (12.5, "12.5"),
    ("  Ρεύμα  ", "Ρεύμα"),
    ("line one\nline two", "line one / line two"),
    ("line one\r\nline two", "line one / line two"),   # Windows
    ("line one\rline two", "line one / line two"),     # old Mac
    *[pytest.param(value, expected, id=name) for name, value, expected in [
        # google docs leaves empty paragraphs inside table cells
        ("empty paragraphs at the end", "ε) Ραντεβού\n\n\n", "ε) Ραντεβού"),
        ("empty paragraph between", "α) one\n\nβ) two", "α) one / β) two"),
        ("empty and blank first", "\n  \nline", "line"),
        ("spaces around the break", "line one  \n  line two", "line one / line two"),
        ("empty windows line", "a\r\n\r\nb", "a / b"),
    ]],
])
def test_cell_to_str(value, expected):
    assert cell_to_str(value) == expected


# what counts as a value and what as a name
@pytest.mark.parametrize("cell, expected", [
    ("180000", True), ("54,5", True), ("2024-03-01", True), ("01/03/2024", True), ("€54", True),
    ("ΗΥ100", False), ("Όνομα", False), ("", False), ("N/A", False),
])
def test_looks_like_value(cell, expected):
    assert tabular.looks_like_value(cell) == expected


N = None
# known limit: in a table of words, a totals row with numbers looks like the first data row
TOTALS_LIMIT = pytest.mark.xfail(strict=True, reason="L29 limit: words table with a numeric totals row")

# (rows, header_rows, expected text)
FORMAT_ROWS = [
    pytest.param([("Ημερομηνία", "Κατηγορία", "Ποσό"), (datetime.datetime(2024, 3, 1), "Ρεύμα", 54),
                  (None, None, None), ("2024-04-01", "Νερό", 12.5)], 0,
                 "Ημερομηνία: 2024-03-01; Κατηγορία: Ρεύμα; Ποσό: 54\n"
                 "Ημερομηνία: 2024-04-01; Κατηγορία: Νερό; Ποσό: 12.5",
                 id="first row is the header, empty rows dropped"),
    pytest.param([[N, N, N, N, N, N, N],
                  [N, "CS-209: English IV, Spring 2026", N, N, N, N, N],
                  [N, N, "TEAM MEMBERS", "Paper", N, "Phases", N],
                  [N, "TEAM #", "NAMES & AM", "Link", "A+B", "C+D", "TOTAL"],
                  [N, 1, "Aggelos Papanikolaou - 5601 (leader)\nEirini Lyroni - 5690",
                   "https://doi.org/10.1145/3635636.3656185", N, N, N]], 0,
                 "CS-209: English IV, Spring 2026\n"
                 "TEAM MEMBERS | Paper | Phases\n"
                 "TEAM #: 1; NAMES & AM: Aggelos Papanikolaou - 5601 (leader) / Eirini Lyroni - 5690; "
                 "Link: https://doi.org/10.1145/3635636.3656185\n"
                 "A+B | C+D | TOTAL",
                 id="title rows above the header are kept as text"),
    pytest.param([("Ημερομηνία", "Κατηγορία", None), ("2024-03-15", None, 20)], 0,
                 "Ημερομηνία: 2024-03-15; col3: 20\nΚατηγορία",
                 id="empty header cell gets a column name"),
    pytest.param([("A", "B", "C"), ("1",)], 0, "A: 1\nB | C", id="short rows do not crash"),
    pytest.param([("Προϊόν", "Ποσότητα", "Τιμή")], 0, "Προϊόν | Ποσότητα | Τιμή", id="only the header"),
    pytest.param([("Ψώνια",), ("γάλα",), ("ψωμί",)], 0, "Ψώνια\nγάλα\nψωμί",
                 id="single column is not a table"),
    pytest.param([], 0, "", id="empty sheet"),
    pytest.param([(None, None), ("", "")], 0, "", id="sheet of empty cells"),
    pytest.param([("Ημερομηνία", "Περιγραφή"), ("2024-03-01", "Ρεύμα", "πληρώθηκε")], 0,
                 "Ημερομηνία: 2024-03-01; Περιγραφή: Ρεύμα; col3: πληρώθηκε",
                 id="wider data row with a date does not become the header"),
    pytest.param([("Οικογενειακά", "έξοδα"), ("", "Κατηγορία", "Ποσό"), ("2024-03-01", "Ρεύμα", "54")], 0,
                 "Οικογενειακά | έξοδα\ncol1: 2024-03-01; Κατηγορία: Ρεύμα; Ποσό: 54",
                 id="climbing stops at the header, not at a title above it"),
    pytest.param([("Οικογενειακά", "έξοδα"), ("", "Κατηγορία", "Ποσό"), ("2024-03-01", "Ρεύμα")], 0,
                 "Οικογενειακά | έξοδα\ncol1: 2024-03-01; Κατηγορία: Ρεύμα\nΠοσό",
                 id="title as full as the header and the data"),
    pytest.param([("Όνομα", "Βαθμός"), ("Μαρία", "9"), ("Νίκος", "απών"), ("Ελένη", "8")], 0,
                 "Όνομα: Μαρία; Βαθμός: 9\nΌνομα: Νίκος; Βαθμός: απών\nΌνομα: Ελένη; Βαθμός: 8",
                 id="data row with only words among the data"),
    pytest.param([("Περιοχή", "2023", "2024"), ("Κρήτη", "άγνωστο", "άγνωστο"), ("Αττική", "120", "150")], 0,
                 "Περιοχή: Κρήτη; 2023: άγνωστο; 2024: άγνωστο\nΠεριοχή: Αττική; 2023: 120; 2024: 150",
                 id="header with years followed by a row of words"),
    pytest.param([("Όνομα", "Μάθημα", "Βαθμός"), ("Χειμερινό", "εξάμηνο"), ("Μαρία", "ΗΥ100", "9")], 0,
                 "Όνομα: Χειμερινό; Μάθημα: εξάμηνο\nΌνομα: Μαρία; Μάθημα: ΗΥ100; Βαθμός: 9",
                 id="two-cell section row does not hide the header above it"),
    pytest.param([("Πωλήσεις",), ("Περιοχή", "2023", "2024"), ("Κρήτη", "120", "150")], 0,
                 "Πωλήσεις\nΠεριοχή: Κρήτη; 2023: 120; 2024: 150",
                 id="one-cell title above a header with years"),
    pytest.param([("", "Μάθημα", "Βαθμός"), ("Χειμερινό",), ("Μαρία", "ΗΥ100", "9")], 0,
                 "col1: Χειμερινό\ncol1: Μαρία; Μάθημα: ΗΥ100; Βαθμός: 9",
                 id="one-cell section row between the header and the data"),
    pytest.param([("Εξάμηνο", "Χειμερινό"), ("Όνομα", "Μάθημα", "Βαθμός"), ("Μαρία", "ΗΥ100", "9")], 0,
                 "Εξάμηνο | Χειμερινό\nΌνομα: Μαρία; Μάθημα: ΗΥ100; Βαθμός: 9",
                 id="subtitle of words above the header stays a subtitle"),
    pytest.param([("Εξάμηνο", "2024"), ("Όνομα", "Μάθημα", "Βαθμός"), ("Μαρία", "ΗΥ100", "9")], 0,
                 "Εξάμηνο | 2024\nΌνομα: Μαρία; Μάθημα: ΗΥ100; Βαθμός: 9",
                 id="subtitle with a year above the header stays a subtitle"),
    pytest.param([("Όνομα", "Επώνυμο"), ("Μαρία", "Παπαδάκη"), ("Νίκος", "Κωστάκης")], 0,
                 "Όνομα: Μαρία; Επώνυμο: Παπαδάκη\nΌνομα: Νίκος; Επώνυμο: Κωστάκης", id="only words"),
    pytest.param([("Μαρία", "9"), ("Νίκος", "8", "άριστα")], 0, "Μαρία | 9\nΝίκος | 8 | άριστα",
                 id="data rows without a header stay plain"),
    pytest.param([("2024", "120", "340"), ("2025", "150", "380")], 0, "2024 | 120 | 340\n2025 | 150 | 380",
                 id="only numbers: no header"),
    pytest.param([("2024", "120", ""), ("2025", "150", "")], 0, "2024 | 120\n2025 | 150",
                 id="only numbers with an empty trailing column (csv lines ending in a comma)"),
    pytest.param([("Πωλήσεις",), ("", "2023", "2024"), ("Κρήτη", "120", "150")], 0,
                 "Πωλήσεις\ncol1: Κρήτη; 2023: 120; 2024: 150",
                 id="pivot table with a title above"),
    pytest.param([("Όνομα", "Επώνυμο", "Τμήμα"), ("", "Παπαδάκη", "ΗΥ"), ("Νίκος", "Κωστάκης", "ΗΥ")], 0,
                 "Επώνυμο: Παπαδάκη; Τμήμα: ΗΥ\nΌνομα: Νίκος; Επώνυμο: Κωστάκης; Τμήμα: ΗΥ",
                 id="words table where a data row lost its first cell"),
    pytest.param([("Περιοχή", "2023", "2024"), ("Κρήτη", "120", "150"), ("", "130", "160")], 0,
                 "Περιοχή: Κρήτη; 2023: 120; 2024: 150\n2023: 130; 2024: 160",
                 id="a data row without its label is still data"),
    pytest.param([("2023", "2024"), ("120", "150")], 1, "2023: 120; 2024: 150",
                 id="header marked by the file is trusted even if numeric"),
    pytest.param([("Όνομα", "Βαθμοί", "Βαθμοί"), ("Όνομα", "Γραπτό", "Προφορικό"), ("Μαρία", "9", "8")], 2,
                 "Όνομα: Μαρία; Βαθμοί - Γραπτό: 9; Βαθμοί - Προφορικό: 8",
                 id="two header rows are joined per column"),
    pytest.param([("Όνομα", "Βαθμός"), ("Μαρία", "9")], 5, "Όνομα | Βαθμός\nΜαρία | 9",
                 id="header_rows larger than the table"),
    # more real shapes
    pytest.param([("Όνομα", "Βαθμός", "Παρατηρήσεις"), ("Μαρία", "9", ""), ("Νίκος", "7", "")], 0,
                 "Όνομα: Μαρία; Βαθμός: 9\nΌνομα: Νίκος; Βαθμός: 7\nΠαρατηρήσεις",
                 id="name of a column without values is kept"),
    pytest.param([("2024-03-01", "ΔΕΗ", "-54,30"), ("2024-03-02", "Μισθός", "1.200,00")], 0,
                 "2024-03-01 | ΔΕΗ | -54,30\n2024-03-02 | Μισθός | 1.200,00",
                 id="bank export without a header"),
    pytest.param([("", "120", "150"), ("2024", "130", "160")], 0, "120 | 150\n2024 | 130 | 160",
                 id="numbers only, empty first cell: not a pivot"),
    pytest.param([("", "120", "150"), ("", "130", "160")], 0, "120 | 150\n130 | 160",
                 id="numbers with an empty first column"),
    pytest.param([("", "Μήνας", "Ποσό"), ("", "Μάρτιος", "54"), ("", "Απρίλιος", "60")], 0,
                 "Μήνας: Μάρτιος; Ποσό: 54\nΜήνας: Απρίλιος; Ποσό: 60",
                 id="table shifted one column right (column A empty)"),
    pytest.param([("Όνομα", "Τμήμα"), ("Μαρία", "ΗΥ"), ("Νίκος", "ΜΑΘ"), ("ΣΥΝΟΛΟ", "2")], 0,
                 "Όνομα: Μαρία; Τμήμα: ΗΥ\nΌνομα: Νίκος; Τμήμα: ΜΑΘ\nΌνομα: ΣΥΝΟΛΟ; Τμήμα: 2",
                 marks=TOTALS_LIMIT, id="words table with a totals row in capitals"),
    pytest.param([("", "Λίστα", "μαθημάτων"), ("Κωδικός", "Τίτλος", "Διδάσκων", "Εξάμηνο"),
                  ("ΗΥ100", "Εισαγωγή", "Παπαδάκης", "Α")], 0,
                 "Λίστα | μαθημάτων\nΚωδικός: ΗΥ100; Τίτλος: Εισαγωγή; Διδάσκων: Παπαδάκης; Εξάμηνο: Α",
                 id="title starting in column B is not a pivot header"),
]


# a table becomes text
@pytest.mark.parametrize("rows, header_rows, expected", FORMAT_ROWS)
def test_format_rows(rows, header_rows, expected):
    if header_rows:
        assert format_rows(rows, header_rows=header_rows) == expected
    else:
        assert format_rows(rows) == expected


# every combination of a pivot table: corner, column names, row labels, cells
PIVOT_PARTS = {
    "corner": {"empty corner": "", "word corner": "Έτος"},
    "columns": {"places across": ["Κρήτη", "Κέρκυρα"], "years across": ["2023", "2024"]},
    "labels": {"places down": ["Ρόδος", "Κως"], "years down": ["2023", "2024"]},
    "cells": {"numbers": [["120", "80"], ["150", "90"]], "words": [["υψηλή", "χαμηλή"], ["μέτρια", "υψηλή"]]},
}


def pivot_cases():
    cases = []
    for parts in itertools.product(*[d.items() for d in PIVOT_PARTS.values()]):
        (cn, corner), (coln, columns), (ln, labels), (celln, cells) = parts
        # all values and an empty corner: it looks exactly like a table of numbers, nobody can tell
        if not corner and (coln, ln, celln) == ("years across", "years down", "numbers"):
            continue
        rows = [(corner, *columns)] + [(label, *row) for label, row in zip(labels, cells)]
        names = [corner or "col1", *columns]
        lines = ["; ".join(f"{n}: {v}" for n, v in zip(names, row) if v) for row in rows[1:]]
        cases.append(pytest.param(rows, "\n".join(lines), id=", ".join([cn, coln, ln, celln])))
    return cases


@pytest.mark.parametrize("rows, expected", pivot_cases())
def test_format_rows_every_pivot_combination(rows, expected):
    assert format_rows(rows) == expected


# an extra value far down the table becomes col3
def test_format_rows_keeps_values_beyond_the_header_far_down_the_table():
    rows = [
        ("Όνομα", "Βαθμός"),
        *[(f"Φοιτητής {i}", "8") for i in range(1, 11)],
        ("Μαρία", "9", "άριστα"),
    ]

    assert format_rows(rows).splitlines()[-1] == "Όνομα: Μαρία; Βαθμός: 9; col3: άριστα"


# csv and tsv files
@pytest.mark.parametrize("name, data, expected", [
    ("greek.csv", "Ημερομηνία;Κατηγορία;Ποσό\n2024-03-01;Ρεύμα;54,5\n".encode("cp1253"),
     "Ημερομηνία: 2024-03-01; Κατηγορία: Ρεύμα; Ποσό: 54,5"),
    ("address.csv", b'date,address,amount\n2024-03-01,"Knossou 5, Heraklion",54\n',
     "date: 2024-03-01; address: Knossou 5, Heraklion; amount: 54"),
    ("table.tsv", b"a\tb\n1\t2\n", "a: 1; b: 2"),
    ("names.csv", b"name\nmaria\ngiorgos\n", "name\nmaria\ngiorgos"),
    ("excel.csv", "Όνομα;Βαθμός\r\nΜαρία;9\r\n".encode("utf-8-sig"), "Όνομα: Μαρία; Βαθμός: 9"),  # with BOM
    ("lf.csv", b'name,comment\r\nmaria,"first line\nsecond line"\r\n',      # Excel, LibreOffice
     "name: maria; comment: first line / second line"),
    ("crlf.csv", b'name,comment\r\nmaria,"first line\r\nsecond line"\r\n',  # Windows editor
     "name: maria; comment: first line / second line"),
    ("bills.csv", "Περιγραφή;Ποσό\n\"Ρεύμα, νερό, τηλέφωνο\";54\n\"Ενοίκιο, κοινόχρηστα\";400\n".encode(),
     "Περιγραφή: Ρεύμα, νερό, τηλέφωνο; Ποσό: 54\nΠεριγραφή: Ενοίκιο, κοινόχρηστα; Ποσό: 400"),
    # pandas index column
    ("pandas.csv", b",city,population\n0,Heraklion,180000\n1,Chania,110000\n",
     "col1: 0; city: Heraklion; population: 180000\ncol1: 1; city: Chania; population: 110000"),
    ("pandas_gap.csv", b",city,population\n0,Heraklion,\n1,Chania,110000\n",
     "col1: 0; city: Heraklion\ncol1: 1; city: Chania; population: 110000"),
    ("ragged.csv", b"name,grade\nmaria,9,excellent\n", "name: maria; grade: 9; col3: excellent"),
])
def test_read_csv(tmp_path, name, data, expected):
    f = tmp_path / name
    f.write_bytes(data)

    assert read_csv(f) == expected


# writes an .xlsx file, sheets is a list of (sheet name, rows)
def make_xlsx(path, sheets):
    book = openpyxl.Workbook()
    book.remove(book.worksheets[0])
    for name, rows in sheets:
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    book.save(path)


# writes an .ods file the same way
def make_ods(path, sheets):
    with pd.ExcelWriter(path, engine="odf") as writer:
        for name, rows in sheets:
            pd.DataFrame(rows).to_excel(writer, sheet_name=name, header=False, index=False)


EXPENSES = [["Ημερομηνία", "Κατηγορία", "Ποσό"], [datetime.datetime(2024, 3, 1), "Ρεύμα", 54]]
INCOME = [["Μήνας", "Ποσό"], ["Μάρτιος", 1200]]
EXPECTED_TWO_SHEETS = (
    "# Sheet: Έξοδα\n"
    "Ημερομηνία: 2024-03-01; Κατηγορία: Ρεύμα; Ποσό: 54\n"
    "\n"
    "# Sheet: Έσοδα\n"
    "Μήνας: Μάρτιος; Ποσό: 1200"
)
EXPECTED_XLS = (
    "# Sheet: Έξοδα\n"
    "Ημερομηνία: 2024-03-01; Κατηγορία: Ρεύμα; Ποσό: 54\n"
    "Ημερομηνία: 2024-04-01; Κατηγορία: Νερό; Ποσό: 12.5\n"
    "\n"
    "# Sheet: Έσοδα\n"
    "Μήνας: Μάρτιος; Ποσό: 1200"
)


# every sheet becomes a block with its name, empty sheets are left out
def test_read_xlsx_all_sheets_and_skips_empty_ones(tmp_path):
    f = tmp_path / "book.xlsx"
    make_xlsx(f, [("Έξοδα", EXPENSES), ("Κενό", []), ("Έσοδα", INCOME)])

    assert read_xlsx(f) == EXPECTED_TWO_SHEETS


# old .xls file: dates and numbers come out like in xlsx
def test_read_xls_dates_numbers_and_sheets():
    assert read_xls(FIXTURES / "expenses.xls") == EXPECTED_XLS


# a line break inside an excel cell becomes " / "
def test_read_xlsx_cell_with_windows_line_break(tmp_path):
    f = tmp_path / "notes.xlsx"
    make_xlsx(f, [("Σημειώσεις", [["Όνομα", "Σχόλιο"], ["Μαρία", "πρώτη γραμμή\r\nδεύτερη γραμμή"]])])

    assert read_xlsx(f) == "# Sheet: Σημειώσεις\nΌνομα: Μαρία; Σχόλιο: πρώτη γραμμή / δεύτερη γραμμή"


# for a formula we take the result excel saved, not the formula
def test_read_xlsx_formula_uses_the_value_excel_saved(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Τιμές"
    sheet.append(["Τεμάχια", "Τιμή", "Σύνολο"])
    sheet.append([2, 5, "=A2*B2"])
    plain = tmp_path / "plain.xlsx"
    book.save(plain)
    # add the result that excel saves next to the formula
    f = tmp_path / "formula.xlsx"
    with zipfile.ZipFile(plain) as src, zipfile.ZipFile(f, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                assert b"<f>A2*B2</f><v></v>" in data
                data = data.replace(b"<f>A2*B2</f><v></v>", b"<f>A2*B2</f><v>10</v>")
            dst.writestr(item, data)

    assert read_xlsx(f) == "# Sheet: Τιμές\nΤεμάχια: 2; Τιμή: 5; Σύνολο: 10"


# parse_file sends xlsx, ods and csv to the table readers
def test_parse_file_uses_the_tabular_readers(tmp_path):
    root = tmp_path / "corpus"
    root.mkdir()
    make_xlsx(root / "book.xlsx", [("Έξοδα", EXPENSES)])
    make_ods(root / "book.ods", [("Έσοδα", INCOME)])
    (root / "data.csv").write_bytes(b"a;b\n1;2\n")

    docs = {}
    for scanned in scan(CorpusConfig(roots=[root])).files:
        docs[scanned.rel_path.name] = parse_file(scanned)

    assert docs["book.xlsx"].raw_text.startswith("# Sheet: Έξοδα")
    assert docs["book.ods"].raw_text.startswith("# Sheet: Έσοδα")
    assert docs["data.csv"].raw_text == "a: 1; b: 2"
    for doc in docs.values():
        assert doc.content_type == "tabular"
        assert doc.parse_status == "ok"


# every merge shape, in every format that has merged cells
# (rows as the user sees them: None = covered by a merge, merges as (top, left, bottom, right))
MERGE_SHAPES = {
    "across": (
        [("Όνομα", "Γραπτό", "Προφορικό"), ("Μαρία", "9", "8"), ("Νίκος", "απαλλαγή", N)], [(2, 1, 2, 2)],
        "Όνομα: Μαρία; Γραπτό: 9; Προφορικό: 8\nΌνομα: Νίκος; Γραπτό: απαλλαγή; Προφορικό: απαλλαγή"),
    "down": (
        [("Κατηγορία", "Μήνας", "Ποσό"), ("Ρεύμα", "Μάρτιος", "54"), (N, "Απρίλιος", "60")], [(1, 0, 2, 0)],
        "Κατηγορία: Ρεύμα; Μήνας: Μάρτιος; Ποσό: 54\nΚατηγορία: Ρεύμα; Μήνας: Απρίλιος; Ποσό: 60"),
    "down in the last column": (
        [("Όνομα", "Γραπτό", "Σχόλιο"), ("Μαρία", "9", "επανεξέταση"), ("Νίκος", "7", N)], [(1, 2, 2, 2)],
        "Όνομα: Μαρία; Γραπτό: 9; Σχόλιο: επανεξέταση\nΌνομα: Νίκος; Γραπτό: 7; Σχόλιο: επανεξέταση"),
    "2x2 block": (
        [("Όνομα", "Γραπτό", "Προφορικό"), ("Μαρία", "απαλλαγή", N), ("Νίκος", N, N)], [(1, 1, 2, 2)],
        "Όνομα: Μαρία; Γραπτό: απαλλαγή; Προφορικό: απαλλαγή\nΌνομα: Νίκος; Γραπτό: απαλλαγή; Προφορικό: απαλλαγή"),
    "empty merged area": (
        [("Όνομα", "Γραπτό", "Προφορικό"), ("Μαρία", "9", "8"), ("Νίκος", N, N)], [(2, 1, 2, 2)],
        "Όνομα: Μαρία; Γραπτό: 9; Προφορικό: 8\nΌνομα: Νίκος"),
    "title over a table of words": (
        [("Κατάλογος", N, N), ("Όνομα", "Επώνυμο", "Τμήμα"), ("Μαρία", "Παπαδάκη", "Φυσική"),
         ("Νίκος", "Λαμπράκης", "Χημεία")], [(0, 0, 0, 2)],
        "Κατάλογος\nΌνομα: Μαρία; Επώνυμο: Παπαδάκη; Τμήμα: Φυσική\nΌνομα: Νίκος; Επώνυμο: Λαμπράκης; Τμήμα: Χημεία"),
    "title over a table of numbers": (
        [("Βαθμοί", N, N), ("Όνομα", "Γραπτό", "Προφορικό"), ("Μαρία", "9", "8")], [(0, 0, 0, 2)],
        "Βαθμοί\nΌνομα: Μαρία; Γραπτό: 9; Προφορικό: 8"),
    "section row between data rows": (
        [("Όνομα", "Γραπτό", "Προφορικό"), ("Μαρία", "9", "8"), ("Επαναληπτική", N, N), ("Νίκος", "7", "6")],
        [(2, 0, 2, 2)],
        "Όνομα: Μαρία; Γραπτό: 9; Προφορικό: 8\nΕπαναληπτική\nΌνομα: Νίκος; Γραπτό: 7; Προφορικό: 6"),
}


def merged_docx(shape, f):
    rows, merges, _ = MERGE_SHAPES[shape]
    d = docx.Document()
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            t.cell(r, c).text = value or ""
    for top, left, bottom, right in merges:
        t.cell(top, left).merge(t.cell(bottom, right)).text = rows[top][left] or ""
    d.save(str(f))
    return read_docx(f)


def merged_pptx(shape, f):
    rows, merges, _ = MERGE_SHAPES[shape]
    p = pptx.Presentation()
    slide = p.slides.add_slide(p.slide_layouts[LAYOUT_BLANK])
    t = slide.shapes.add_table(len(rows), len(rows[0]), PptxInches(1), PptxInches(1), PptxInches(6),
                               PptxInches(2)).table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            t.cell(r, c).text = value or ""
    for top, left, bottom, right in merges:
        t.cell(top, left).merge(t.cell(bottom, right))
        t.cell(top, left).text = rows[top][left] or ""
    p.save(str(f))
    return read_pptx(f).removeprefix("# Slide 1\n\n")


def merged_xlsx(shape, f):
    rows, merges, _ = MERGE_SHAPES[shape]
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    for top, left, bottom, right in merges:
        sheet.merge_cells(start_row=top + 1, start_column=left + 1, end_row=bottom + 1, end_column=right + 1)
    book.save(f)
    return read_xlsx(f).removeprefix("# Sheet: Sheet\n")


# the merge shape as an odf table: the first cell of a merge says its size, the cells it hides are covered
def merged_odf_table(shape):
    rows, merges, _ = MERGE_SHAPES[shape]
    origins = {(t, l): (b - t + 1, r - l + 1) for t, l, b, r in merges}
    covered = {(r, c) for t, l, b, rr in merges for r in range(t, b + 1) for c in range(l, rr + 1)} - set(origins)
    table = OdfTable(name="Φύλλο")
    for r, row in enumerate(rows):
        tr = TableRow()
        for c, value in enumerate(row):
            if (r, c) in covered:
                tr.addElement(CoveredTableCell())
                continue
            down, across = origins.get((r, c), (1, 1))
            cell = TableCell(valuetype="string", numberrowsspanned=down, numbercolumnsspanned=across)
            if value:
                cell.addElement(OdfP(text=value))
            tr.addElement(cell)
        table.addElement(tr)
    return table


def merged_ods(shape, f):
    doc = OpenDocumentSpreadsheet()
    doc.spreadsheet.addElement(merged_odf_table(shape))  # pyright: ignore[reportAttributeAccessIssue] (odfpy adds it at runtime)
    doc.save(str(f))
    return read_ods(f).removeprefix("# Sheet: Φύλλο\n")


# the same table inside a writer document
def merged_odt(shape, f):
    doc = OpenDocumentText()
    doc.text.addElement(merged_odf_table(shape))  # pyright: ignore[reportAttributeAccessIssue] (odfpy adds it at runtime)
    doc.save(str(f))
    return read_odf(f)


# nothing writes xls today: one sheet per shape, made once with xlwt
def merged_xls(shape, f):
    sheets = {}
    for block in read_xls(FIXTURES / "merged.xls").split("\n\n"):
        name, text = block.removeprefix("# Sheet: ").split("\n", 1)
        sheets[name] = text
    return sheets[shape]


# the same table on a slide of an impress presentation
def merged_odp(shape, f):
    doc = OpenDocumentPresentation()
    frame = Frame()
    frame.addElement(merged_odf_table(shape))
    page = Page(masterpagename="Default")
    page.addElement(frame)
    doc.presentation.addElement(page)  # pyright: ignore[reportAttributeAccessIssue] (odfpy adds it at runtime)
    doc.save(str(f))
    return read_odf(f).removeprefix("# Slide 1\n\n")


MERGE_FORMATS = {"docx": merged_docx, "pptx": merged_pptx, "xlsx": merged_xlsx, "xls": merged_xls, "ods": merged_ods,
                 "odt": merged_odt, "odp": merged_odp}


@pytest.mark.parametrize("fmt, shape", [pytest.param(fmt, shape, id=f"{fmt}, {shape}")
                                        for fmt, shape in itertools.product(MERGE_FORMATS, MERGE_SHAPES)])
def test_merged_cells(tmp_path, fmt, shape):
    assert MERGE_FORMATS[fmt](shape, tmp_path / f"merged.{fmt}") == MERGE_SHAPES[shape][2]


# xlsx merges from the sheet xml: tags cut between two chunks, and tags with a prefix
@pytest.mark.parametrize("chunk, prefix", [(7, ""), (1024 * 1024, "x:")], ids=["tiny chunks", "x: prefix"])
def test_read_xlsx_merges_from_the_xml(tmp_path, monkeypatch, chunk, prefix):
    merged_xlsx("across", tmp_path / "plain.xlsx")
    f = tmp_path / "merged.xlsx"
    with zipfile.ZipFile(tmp_path / "plain.xlsx") as zin, zipfile.ZipFile(f, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                ns = b'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
                data = data.replace(ns, ns + b" " + ns.replace(b"xmlns", b"xmlns:x"), 1)
                data = data.replace(b"<mergeCell ", f"<{prefix}mergeCell ".encode())
            zout.writestr(item, data)
    monkeypatch.setattr(tabular, "MB", chunk)

    assert read_xlsx(f) == "# Sheet: Sheet\n" + MERGE_SHAPES["across"][2]


# an xlsx with one sheet per (name, rows, links); links maps a cell like "B2" to its address
def linked_xlsx(path, sheets):
    book = openpyxl.Workbook()
    book.remove(book.worksheets[0])
    for name, rows, links in sheets:
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
        for cell, address in links.items():
            if address.startswith("#"):
                sheet[cell].hyperlink = Hyperlink(ref=cell, location=address[1:])
            else:
                sheet[cell].hyperlink = address
    book.save(path)
    return read_xlsx(path)


COURSES = [["Μάθημα", "Σελίδα"], ["ΗΥ-252", "Ιστοσελίδα"]]

# (sheets, text)
XLSX_LINK_CASES = {
    "cell with a link": (
        [("Μαθήματα", COURSES, {"B2": "https://www.csd.uoc.gr/~hy252"})],
        "# Sheet: Μαθήματα\nΜάθημα: ΗΥ-252; Σελίδα: Ιστοσελίδα (https://www.csd.uoc.gr/~hy252)"),
    "email link": (
        [("Επαφές", [["Όνομα", "Email"], ["Γραμματεία", "Γράψτε μας"]], {"B2": "mailto:secretary@csd.uoc.gr"})],
        "# Sheet: Επαφές\nΌνομα: Γραμματεία; Email: Γράψτε μας (mailto:secretary@csd.uoc.gr)"),
    "number with a link": (
        [("Βαθμοί", [["Μάθημα", "Βαθμός"], ["ΗΥ-252", 9]], {"B2": "https://elearn.uoc.gr/grades"})],
        "# Sheet: Βαθμοί\nΜάθημα: ΗΥ-252; Βαθμός: 9 (https://elearn.uoc.gr/grades)"),
    "each sheet has its own links": (
        [("Α", [["Μάθημα", "Σελίδα"], ["ΗΥ-100", "Πρώτη"]], {"B2": "https://a.gr"}),
         ("Β", [["Μάθημα", "Σελίδα"], ["ΗΥ-200", "Δεύτερη"]], {"B2": "https://b.gr"})],
        "# Sheet: Α\nΜάθημα: ΗΥ-100; Σελίδα: Πρώτη (https://a.gr)\n\n# Sheet: Β\nΜάθημα: ΗΥ-200; Σελίδα: Δεύτερη (https://b.gr)"),
    "two links in one sheet": (
        [("Μαθήματα", [["Μάθημα", "Σελίδα"], ["ΗΥ-100", "Πρώτη"], ["ΗΥ-200", "Δεύτερη"]],
          {"B2": "https://a.gr", "B3": "https://b.gr"})],
        "# Sheet: Μαθήματα\nΜάθημα: ΗΥ-100; Σελίδα: Πρώτη (https://a.gr)\nΜάθημα: ΗΥ-200; Σελίδα: Δεύτερη (https://b.gr)"),
}


@pytest.mark.parametrize("sheets, expected", XLSX_LINK_CASES.values(), ids=XLSX_LINK_CASES.keys())
def test_read_xlsx_links(tmp_path, sheets, expected):
    assert linked_xlsx(tmp_path / "links.xlsx", sheets) == expected


# (sheets, text) links that add nothing
XLSX_LINK_KEPT = {
    "link that is the text of the cell is not written twice": (
        [("Μαθήματα", [["Μάθημα", "Σελίδα"], ["ΗΥ-252", "https://www.csd.uoc.gr/~hy252"]], {"B2": "https://www.csd.uoc.gr/~hy252"})],
        "# Sheet: Μαθήματα\nΜάθημα: ΗΥ-252; Σελίδα: https://www.csd.uoc.gr/~hy252"),
    "link to a place in the workbook is left out": (
        [("Μαθήματα", COURSES, {"B2": "#Βαθμοί!A1"})],
        "# Sheet: Μαθήματα\nΜάθημα: ΗΥ-252; Σελίδα: Ιστοσελίδα"),
}


@pytest.mark.parametrize("sheets, expected", XLSX_LINK_KEPT.values(), ids=XLSX_LINK_KEPT.keys())
def test_read_xlsx_links_that_add_nothing(tmp_path, sheets, expected):
    assert linked_xlsx(tmp_path / "links.xlsx", sheets) == expected


# a date with a link is written like any date
def test_read_xlsx_date_with_a_link(tmp_path):
    rows = [["Ημερομηνία", "Θέμα"], [datetime.datetime(2024, 3, 1), "Εξέταση"]]
    text = linked_xlsx(tmp_path / "dates.xlsx", [("Πρόγραμμα", rows, {"A2": "https://www.csd.uoc.gr/exams"})])
    assert text == "# Sheet: Πρόγραμμα\nΗμερομηνία: 2024-03-01 (https://www.csd.uoc.gr/exams); Θέμα: Εξέταση"


# a merged cell with a link has the link in every cell it covers
def test_read_xlsx_merged_cell_with_a_link(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Βαθμοί"
    for row in [["Όνομα", "Πρόοδος", "Τελικό"], ["Μαρία", "Προφίλ", None]]:
        sheet.append(row)
    sheet["B2"].hyperlink = "https://elearn.uoc.gr/maria"
    sheet.merge_cells("B2:C2")
    f = tmp_path / "merged.xlsx"
    book.save(f)
    assert read_xlsx(f) == ("# Sheet: Βαθμοί\nΌνομα: Μαρία; Πρόοδος: Προφίλ (https://elearn.uoc.gr/maria); "
                            "Τελικό: Προφίλ (https://elearn.uoc.gr/maria)")


# a link on a cell outside the data of the sheet is left out, the file still reads
def test_read_xlsx_link_outside_the_data(tmp_path):
    linked_xlsx(tmp_path / "plain.xlsx", [("Μαθήματα", COURSES, {"B2": "https://www.csd.uoc.gr/~hy252"})])
    f = tmp_path / "outside.xlsx"
    with zipfile.ZipFile(tmp_path / "plain.xlsx") as zin, zipfile.ZipFile(f, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'ref="B2"', b'ref="H40"')
            zout.writestr(item, data)
    assert read_xlsx(f) == "# Sheet: Μαθήματα\nΜάθημα: ΗΥ-252; Σελίδα: Ιστοσελίδα"


# a sheet without links reads as before
def test_read_xlsx_without_links(tmp_path):
    assert linked_xlsx(tmp_path / "plain.xlsx", [("Μαθήματα", COURSES, {})]) == "# Sheet: Μαθήματα\nΜάθημα: ΗΥ-252; Σελίδα: Ιστοσελίδα"


# links from the sheet xml: tags cut between two chunks, and tags with a prefix
@pytest.mark.parametrize("chunk, prefix", [(7, ""), (1024 * 1024, "x:")], ids=["tiny chunks", "x: prefix"])
def test_read_xlsx_links_from_the_xml(tmp_path, monkeypatch, chunk, prefix):
    linked_xlsx(tmp_path / "plain.xlsx", [("Μαθήματα", COURSES, {"B2": "https://www.csd.uoc.gr/~hy252"})])
    f = tmp_path / "linked.xlsx"
    with zipfile.ZipFile(tmp_path / "plain.xlsx") as zin, zipfile.ZipFile(f, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                ns = b'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
                data = data.replace(ns, ns + b" " + ns.replace(b"xmlns", b"xmlns:x"), 1)
                data = data.replace(b"<hyperlink ", f"<{prefix}hyperlink ".encode())
            zout.writestr(item, data)
    monkeypatch.setattr(tabular, "MB", chunk)
    assert read_xlsx(f) == "# Sheet: Μαθήματα\nΜάθημα: ΗΥ-252; Σελίδα: Ιστοσελίδα (https://www.csd.uoc.gr/~hy252)"


# the value of a merge goes into every cell it covers
@pytest.mark.parametrize("rows, merges, expected", [
    ([("a", "b"), ("c", "d")], [], [["a", "b"], ["c", "d"]]),
    ([("a", None, None)], [(0, 0, 0, 2)], [["a", "a", "a"]]),
    ([("a",), ("b",)], [(0, 0, 0, 2)], [["a", "a", "a"], ["b"]]),
    ([("a", "b"), (None, "c")], [(0, 0, 3, 0)], [["a", "b"], ["a", "c"]]),
    ([(54, None), (None, None)], [(0, 0, 1, 1)], [[54, 54], [54, 54]]),
    ([(None, None), ("x", None)], [(0, 0, 0, 1), (1, 0, 1, 1)], [[None, None], ["x", "x"]]),
    ([("a",)], [(3, 0, 4, 1)], [["a"]]),
    ([("a",)], [(0, 2, 0, 3)], [["a"]]),
], ids=["no merges", "across", "past the end of a short row", "past the last row", "raw value kept",
        "empty merge", "merge below the table", "merge right of a short row"])
def test_fill_merged(rows, merges, expected):
    assert fill_merged(rows, merges) == expected


# a name merged over two columns, under it the names of the two columns
def test_merged_header_over_two_rows():
    rows = fill_merged([("Όνομα", "Βαθμοί", None), (None, "Γραπτό", "Προφορικό"), ("Μαρία", "9", "8")],
                               [(0, 0, 1, 0), (0, 1, 0, 2)])

    assert format_rows(rows, header_rows=2) == "Όνομα: Μαρία; Βαθμοί - Γραπτό: 9; Βαθμοί - Προφορικό: 8"


# word style name to heading level (0 means not a heading)
@pytest.mark.parametrize("style, level", [
    ("Title", 1), ("Heading 1", 1), ("Heading 2", 2), ("Heading 4", 4), ("Heading 9", 9),
    ("Normal", 0), ("Caption", 0), ("Heading", 0), ("Heading Char", 0),
])
def test_heading_level(style, level):
    assert heading_level(style) == level


# a word document with everything read_docx has to handle
def make_docx(path):
    image = path.parent / "img.png"
    Image.new("RGB", (60, 30), "white").save(image)

    doc = docx.Document()
    doc.add_heading("Εισαγωγή", level=1)
    doc.add_paragraph("Πρώτη παράγραφος.\nΜε δεύτερη γραμμή.")
    doc.add_paragraph("")
    doc.add_paragraph("   ")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Όνομα", "Βαθμός"
    table.cell(1, 0).text, table.cell(1, 1).text = "Μαρία", "9"
    doc.add_heading("Μέθοδος", level=2)
    doc.add_picture(str(image), width=Inches(1))
    doc.inline_shapes[-1]._inline.docPr.set("descr", "Αρχιτεκτονική του συστήματος")
    doc.add_picture(str(image), width=Inches(1))
    doc.inline_shapes[-1]._inline.docPr.set("descr", "")   # empty alt text
    doc.add_paragraph("Εικόνα 1: Η ροή των δεδομένων", style="Caption")
    doc.add_heading("Λεπτομέρειες", level=4)
    doc.add_paragraph("Τελευταία παράγραφος.")
    doc.save(str(path))


EXPECTED_DOCX = (
    "# Εισαγωγή\n"
    "\n"
    "Πρώτη παράγραφος.\nΜε δεύτερη γραμμή.\n"
    "\n"
    "Όνομα: Μαρία; Βαθμός: 9\n"
    "\n"
    "## Μέθοδος\n"
    "\n"
    "[Image: Αρχιτεκτονική του συστήματος]\n"
    "\n"
    "Εικόνα 1: Η ροή των δεδομένων\n"
    "\n"
    "#### Λεπτομέρειες\n"
    "\n"
    "Τελευταία παράγραφος."
)

LAYOUT_TITLE_AND_CONTENT = 1
LAYOUT_TITLE_ONLY = 5
LAYOUT_BLANK = 6


# sets the alt text of a picture where powerpoint keeps it
def set_alt_text(shape, text):
    shape._element.xpath("./*/p:cNvPr")[0].set("descr", text)


# writes the speaker notes of a slide
def set_notes(slide, text):
    frame = slide.notes_slide.notes_text_frame
    assert frame is not None
    frame.text = text


# a presentation with everything read_pptx has to handle
def make_pptx(path):
    image = path.parent / "img.png"
    Image.new("RGB", (60, 30), "white").save(image)
    inch = PptxInches(1)

    presentation = pptx.Presentation()

    slide = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_TITLE_AND_CONTENT])
    title = slide.shapes.title
    assert title is not None
    title.text = "  Εισαγωγή "
    body = slide.placeholders[1]
    assert isinstance(body, PptxShape)
    body.text_frame.text = "Πρώτη κουκκίδα"
    body.text_frame.add_paragraph().text = "Δεύτερη κουκκίδα"
    slide.shapes.add_textbox(inch, PptxInches(5), inch, inch)          # empty text box
    set_notes(slide, "να πω για το dataset")

    slide = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    table = slide.shapes.add_table(2, 2, inch, inch, PptxInches(4), inch).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Όνομα", "Βαθμός"
    table.cell(1, 0).text, table.cell(1, 1).text = "Μαρία", "9"
    set_alt_text(slide.shapes.add_picture(str(image), inch, PptxInches(3)), "Διάγραμμα")
    set_alt_text(slide.shapes.add_picture(str(image), PptxInches(3), PptxInches(3)), "")
    outer = slide.shapes.add_group_shape()
    box = outer.shapes.add_textbox(PptxInches(5), PptxInches(5), inch, inch)
    box.text_frame.text = "Κείμενο μέσα σε ομάδα"
    inner = outer.shapes.add_group_shape()
    box = inner.shapes.add_textbox(PptxInches(6), PptxInches(6), inch, inch)
    box.text_frame.text = "Ομάδα μέσα σε ομάδα"

    slide = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_TITLE_ONLY])
    set_notes(slide, "   ")

    presentation.save(str(path))


EXPECTED_PPTX = (
    "# Slide 1: Εισαγωγή\n"
    "\n"
    "Πρώτη κουκκίδα\nΔεύτερη κουκκίδα\n"
    "\n"
    "Notes: να πω για το dataset\n"
    "\n"
    "# Slide 2\n"
    "\n"
    "Όνομα: Μαρία; Βαθμός: 9\n"
    "\n"
    "[Image: Διάγραμμα]\n"
    "\n"
    "Κείμενο μέσα σε ομάδα\n"
    "\n"
    "Ομάδα μέσα σε ομάδα\n"
    "\n"
    "# Slide 3"
)

# not imported from office.py, so a typo there is caught
DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
PPTX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
VARIANT_CONTENT_TYPES = {
    ".docm": "application/vnd.ms-word.document.macroEnabled.main+xml",
    ".dotx": "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
    ".dotm": "application/vnd.ms-word.template.macroEnabledTemplate.main+xml",
    ".pptm": "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
    ".potx": "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
    ".potm": "application/vnd.ms-powerpoint.template.macroEnabled.main+xml",
    ".ppsx": "application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml",
    ".ppsm": "application/vnd.ms-powerpoint.slideshow.macroEnabled.main+xml",
}
WORD_VARIANTS = [".docm", ".dotx", ".dotm"]
SLIDE_VARIANTS_PPTX_REFUSES = [".potx", ".potm", ".ppsx", ".ppsm"]


# a variant is the same file with another label in [Content_Types].xml
def make_variant(path):
    is_word = path.suffix in WORD_VARIANTS
    base = path.with_name("base" + (".docx" if is_word else ".pptx"))
    (make_docx if is_word else make_pptx)(base)
    main = DOCX_CONTENT_TYPE if is_word else PPTX_CONTENT_TYPE
    with zipfile.ZipFile(base) as src, zipfile.ZipFile(path, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(main.encode(), VARIANT_CONTENT_TYPES[path.suffix].encode())
            dst.writestr(item, data)
    base.unlink()


# (extension, maker, reader, expected text)
READERS = [
    (".docx", make_docx, read_docx, EXPECTED_DOCX),
    (".docx", make_docx, read_word_variants, EXPECTED_DOCX),   # also reads a plain docx
    *[(ext, make_variant, read_word_variants, EXPECTED_DOCX) for ext in WORD_VARIANTS],
    (".pptx", make_pptx, read_pptx, EXPECTED_PPTX),
    (".pptm", make_variant, read_pptx, EXPECTED_PPTX),         # python-pptx accepts .pptm itself
    *[(ext, make_variant, read_slide_variants, EXPECTED_PPTX) for ext in SLIDE_VARIANTS_PPTX_REFUSES],
]


# each reader gives exactly the expected text for its file
@pytest.mark.parametrize("ext, make, reader, expected", READERS,
                         ids=[f"{ext}-{reader.__name__}" for ext, _, reader, _ in READERS])
def test_reader_gives_the_expected_text(tmp_path, ext, make, reader, expected):
    f = tmp_path / f"doc{ext}"
    make(f)

    assert reader(f) == expected


# the normal readers refuse the variants
@pytest.mark.parametrize("ext, reader, message", [
    *[(ext, read_docx, "not a Word file") for ext in WORD_VARIANTS],
    *[(ext, read_pptx, "not a PowerPoint file") for ext in SLIDE_VARIANTS_PPTX_REFUSES],
])
def test_main_reader_refuses_the_variants(tmp_path, ext, reader, message):
    f = tmp_path / f"doc{ext}"
    make_variant(f)

    with pytest.raises(ValueError, match=message):
        reader(f)


# the file on disk is not changed
@pytest.mark.parametrize("ext, reader", [(".docm", read_word_variants), (".ppsx", read_slide_variants)])
def test_variant_reader_leaves_the_file_on_disk_unchanged(tmp_path, ext, reader):
    f = tmp_path / f"doc{ext}"
    make_variant(f)
    before = f.read_bytes()

    reader(f)

    assert f.read_bytes() == before


# an empty word or powerpoint file gives empty text
@pytest.mark.parametrize("name, make, reader", [
    ("empty.docx", lambda f: docx.Document().save(str(f)), read_docx),
    ("empty.pptx", lambda f: pptx.Presentation().save(str(f)), read_pptx),
])
def test_empty_document_gives_empty_text(tmp_path, name, make, reader):
    f = tmp_path / name
    make(f)

    assert reader(f) == ""


# a picture gives [Image: alt text]
@pytest.mark.parametrize("alt, expected", [("", []), ("Χάρτης", ["[Image: Χάρτης]"])])
def test_format_shape_picture(tmp_path, alt, expected):
    image = tmp_path / "img.png"
    Image.new("RGB", (60, 30), "white").save(image)
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    picture = slide.shapes.add_picture(str(image), PptxInches(0), PptxInches(0))
    set_alt_text(picture, alt)

    assert format_shape(picture) == expected


RTF_GREEK_CP1253 = (
    "{\\rtf1\\ansi\\ansicpg1253\\deff0{\\fonttbl{\\f0 Arial;}}\r\n"
    "\\f0 \\'c3\\'e5\\'e9\\'e1 \\'f3\\'ef\\'f5\\par\r\n"
    "\\par\r\n"
    "Second paragraph\\line same paragraph, new line\\par\r\n"
    "}"
)
EXPECTED_RTF = "Γεια σου\n\nSecond paragraph\nsame paragraph, new line"


# (file name, maker, content type, expected text)
DOCUMENTS = [
    ("doc.docx", make_docx, "prose", EXPECTED_DOCX),
    *[(f"doc{ext}", make_variant, "prose", EXPECTED_DOCX) for ext in WORD_VARIANTS],
    ("deck.pptx", make_pptx, "prose", EXPECTED_PPTX),
    *[(f"deck{ext}", make_variant, "prose", EXPECTED_PPTX)
      for ext in [".pptm", *SLIDE_VARIANTS_PPTX_REFUSES]],
    ("doc.rtf", lambda f: f.write_bytes(RTF_GREEK_CP1253.encode("ascii")), "prose", EXPECTED_RTF),
    ("expenses.xls", lambda f: f.write_bytes((FIXTURES / "expenses.xls").read_bytes()), "tabular",
     EXPECTED_XLS),
    ("notes.markdown", lambda f: f.write_bytes("# Σημειώσεις\n\nΚείμενο.".encode()), "prose",
     "# Σημειώσεις\n\nΚείμενο."),
]


# every document format through parse_file
@pytest.mark.parametrize("name, make, content_type, expected", DOCUMENTS, ids=[d[0] for d in DOCUMENTS])
def test_parse_file_reads_the_document(tmp_path, name, make, content_type, expected):
    doc = parse_one(tmp_path, name, built(tmp_path, name, make))

    assert doc.parse_status == "ok"
    assert doc.content_type == content_type
    assert doc.raw_text == expected


ZIP_FORMATS = [".docx", *WORD_VARIANTS, ".pptx", ".pptm", *SLIDE_VARIANTS_PPTX_REFUSES, ".xlsx", ".xls", ".ods"]


# a broken file of any format is failed
@pytest.mark.parametrize("name, data", [
    *[(f"broken{ext}", b"this is not a zip file") for ext in ZIP_FORMATS],
    ("broken.rtf", b"this is not an rtf file"),
    ("empty.rtf", b""),
])
def test_parse_file_marks_a_broken_file_as_failed(tmp_path, name, data):
    doc = parse_one(tmp_path, name, data)

    assert doc.parse_status == "failed"
    assert doc.content_type == ("tabular" if name.endswith((".xlsx", ".xls", ".ods")) else "prose")


# an rtf with no text is empty, not failed
def test_parse_file_valid_but_empty_rtf_is_empty(tmp_path):
    assert parse_one(tmp_path, "doc.rtf", rb"{\rtf1\ansi}").parse_status == "empty"


# each extension points to the right reader in PARSERS
@pytest.mark.parametrize("ext, reader", [
    (".docx", read_docx),
    *[(ext, read_word_variants) for ext in WORD_VARIANTS],
    (".pptx", read_pptx),
    (".pptm", read_pptx),
    *[(ext, read_slide_variants) for ext in SLIDE_VARIANTS_PPTX_REFUSES],
    (".rtf", read_rtf),
])
def test_extension_is_registered(ext, reader):
    assert PARSERS.get(ext) is reader


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# a table inside a table cell (common in forms)
def test_read_docx_table_inside_a_table_cell(tmp_path):
    d = docx.Document()
    outer = d.add_table(rows=1, cols=2)
    outer.cell(0, 0).text = "Στοιχεία"
    inner = outer.cell(0, 1).add_table(rows=2, cols=2)
    inner.cell(0, 0).text, inner.cell(0, 1).text = "Τηλέφωνο", "Email"
    inner.cell(1, 0).text, inner.cell(1, 1).text = "2810123456", "maria@uoc.gr"
    f = tmp_path / "nested.docx"
    d.save(str(f))

    text = read_docx(f)

    assert "2810123456" in text
    assert "maria@uoc.gr" in text


# word saves a text box twice, we want it once
def test_read_docx_text_box_appears_once(tmp_path):
    d = docx.Document()
    d.add_paragraph("Πριν")
    box = "<w:txbxContent><w:p><w:r><w:t>Κείμενο σε πλαίσιο</w:t></w:r></w:p></w:txbxContent>"
    d.add_paragraph()._p.append(parse_xml(
        f'<w:r xmlns:w="{W_NS}"'
        ' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
        ' xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"'
        ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        ' xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
        ' xmlns:v="urn:schemas-microsoft-com:vml">'
        '<mc:AlternateContent>'
        '<mc:Choice Requires="wps"><w:drawing><wp:anchor><wp:docPr id="5" name="Text Box 5"/>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f'<wps:wsp><wps:txbx>{box}</wps:txbx><wps:bodyPr/></wps:wsp>'
        '</a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>'
        f'<mc:Fallback><w:pict><v:shape><v:textbox>{box}</v:textbox></v:shape></w:pict></mc:Fallback>'
        '</mc:AlternateContent></w:r>'
    ))
    d.add_paragraph("Μετά")
    f = tmp_path / "textbox.docx"
    d.save(str(f))

    text = read_docx(f)

    assert text.count("Κείμενο σε πλαίσιο") == 1
    assert text.index("Πριν") < text.index("Κείμενο σε πλαίσιο") < text.index("Μετά")


# a document with one paragraph and some raw word xml
def docx_with_xml(tmp_path, first_text, xml_parts, last_text=None, inside=True):
    d = docx.Document()
    p = d.add_paragraph(first_text)
    for xml in xml_parts:
        # add the w: namespace to the first tag
        tag_end = xml.index(" ") if " " in xml.split(">")[0] else xml.index(">")
        element = parse_xml(xml[:tag_end] + f' xmlns:w="{W_NS}"' + xml[tag_end:])
        if inside:
            p._p.append(element)
        else:
            p._p.addnext(element)
    if last_text is not None:
        p.add_run(last_text)
    f = tmp_path / "doc.docx"
    d.save(str(f))
    return f


# text inside content controls, tracked changes and links
@pytest.mark.parametrize("first, xml_parts, last, inside, expected", [
    pytest.param("Αίτηση", ['<w:sdt><w:sdtPr/><w:sdtContent>'
                            '<w:p><w:r><w:t>Όνομα φοιτητή: Μαρία Παπαδάκη</w:t></w:r></w:p>'
                            '</w:sdtContent></w:sdt>'], None, False,
                 "Αίτηση\n\nΌνομα φοιτητή: Μαρία Παπαδάκη",
                 id="content control around paragraphs"),
    pytest.param("Ημερομηνία: ", ['<w:sdt><w:sdtPr/><w:sdtContent>'
                                  '<w:r><w:t>15/06/2026</w:t></w:r></w:sdtContent></w:sdt>'], None, True,
                 "Ημερομηνία: 15/06/2026",
                 id="content control inside a paragraph"),
    pytest.param("Η προθεσμία είναι ",
                 ['<w:del w:id="1" w:author="A" w:date="2026-01-01T00:00:00Z">'
                  '<w:r><w:delText>η Δευτέρα</w:delText></w:r></w:del>',
                  '<w:ins w:id="2" w:author="A" w:date="2026-01-01T00:00:00Z">'
                  '<w:r><w:t>η Τρίτη</w:t></w:r></w:ins>'], ".", True,
                 "Η προθεσμία είναι η Τρίτη.",
                 id="tracked changes give the current text"),
    pytest.param("Δες ", ['<w:hyperlink w:anchor="top"><w:r><w:t>τον οδηγό</w:t></w:r></w:hyperlink>'],
                 " για λεπτομέρειες.", True, "Δες τον οδηγό για λεπτομέρειες.", id="hyperlink text"),
    pytest.param("α", ['<w:r><w:br w:type="page"/><w:t>β</w:t></w:r>'], None, True, "α\nβ", id="page break inside a paragraph"),
    pytest.param("e", ['<w:r><w:noBreakHyphen/><w:t>mail soft</w:t><w:softHyphen/><w:t>ware</w:t></w:r>'], None, True,
                 "e-mail software", id="non-breaking and soft hyphen"),
    pytest.param("α", ['<w:ins w:id="5" w:author="A" w:date="2026-01-01T00:00:00Z"><w:r><w:tab/><w:t>β</w:t><w:br/><w:t>γ</w:t></w:r></w:ins>'], None, True,
                 "α\tβ\nγ", id="tab and line break inside an insertion"),
    pytest.param("Πρώτα ", ['<w:moveFrom w:id="6" w:author="A" w:date="2026-01-01T00:00:00Z"><w:r><w:t>παλιά θέση </w:t></w:r></w:moveFrom>',
                            '<w:moveTo w:id="7" w:author="A" w:date="2026-01-01T00:00:00Z"><w:r><w:t>νέα θέση</w:t></w:r></w:moveTo>'], None, True,
                 "Πρώτα νέα θέση", id="moved text appears once"),
    pytest.param("Συνάντηση στο ", ['<w:smartTag w:uri="urn:x" w:element="place"><w:r><w:t>Ηράκλειο</w:t></w:r>'
                                    '</w:smartTag>'], None, True, "Συνάντηση στο Ηράκλειο", id="smart tag"),
    pytest.param("Σελίδα ", ['<w:fldSimple w:instr=" PAGE "><w:r><w:t>3</w:t></w:r></w:fldSimple>'], None, True,
                 "Σελίδα 3", id="simple field shows its result"),
    pytest.param("Βλ. ", ['<w:r><w:fldChar w:fldCharType="begin"/></w:r>',
                          '<w:r><w:instrText xml:space="preserve"> REF _Ref1 \\h </w:instrText></w:r>',
                          '<w:r><w:fldChar w:fldCharType="separate"/></w:r>', '<w:r><w:t>Πίνακας 2</w:t></w:r>',
                          '<w:r><w:fldChar w:fldCharType="end"/></w:r>'], None, True, "Βλ. Πίνακας 2",
                 id="field result yes, field code no"),
    pytest.param("Κείμενο", ['<w:commentRangeStart w:id="0"/>', '<w:r><w:t xml:space="preserve"> με σχόλιο</w:t></w:r>',
                             '<w:commentRangeEnd w:id="0"/>', '<w:r><w:commentReference w:id="0"/></w:r>'], None, True,
                 "Κείμενο με σχόλιο", id="comment marks are not text"),
    pytest.param("Κείμενο", [f'<w:p><w:pPr><w:rPr><w:del w:id="3" w:author="A" w:date="2026-01-01T00:00:00Z"/></w:rPr></w:pPr><w:del w:id="4" w:author="A" w:date="2026-01-01T00:00:00Z">'
                             '<w:r><w:delText>σβησμένη παράγραφος</w:delText></w:r></w:del></w:p>'], None, False,
                 "Κείμενο", id="deleted paragraph"),
    pytest.param("Κείμενο", ['<w:sdt><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Table of Contents"/>'
                             '<w:docPartUnique/></w:docPartObj></w:sdtPr><w:sdtContent><w:p><w:r><w:t>Εισαγωγή 1</w:t>'
                             '</w:r></w:p></w:sdtContent></w:sdt>'], None, False, "Κείμενο", id="table of contents"),
    pytest.param("Βαθμοί", ['<w:sdt><w:sdtPr/><w:sdtContent><w:tbl><w:tblPr/><w:tblGrid><w:gridCol/><w:gridCol/></w:tblGrid><w:tr><w:tc><w:p><w:r><w:t>Όνομα</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Βαθμός</w:t></w:r></w:p></w:tc></w:tr><w:tr><w:tc><w:p><w:r><w:t>Μαρία</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>9</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:sdtContent></w:sdt>'], None, False,
                 "Βαθμοί\n\nΌνομα: Μαρία; Βαθμός: 9", id="content control around a table"),
    pytest.param("Κείμενο", ['<w:r xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>Πλαίσιο</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>'],
                 None, True, "Πλαίσιο\n\nΚείμενο", id="text box of old word (vml only)"),
    pytest.param("", ['<w:r xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>Έξω</w:t></w:r><w:r><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>Μέσα</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>'],
                 None, True, "Μέσα\n\nΈξω", id="text box inside a text box"),
    pytest.param("", ['<w:r xmlns:v="urn:schemas-microsoft-com:vml" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:drawing><wp:inline><wp:docPr id="1" name="P" descr="Λογότυπο"/></wp:inline></w:drawing></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>'],
                 None, True, "[Image: Λογότυπο]", id="picture inside a text box"),
    pytest.param("α", ['<w:r><w:t/><w:t>β</w:t></w:r>'], None, True, "αβ", id="empty w:t"),
    pytest.param("Κείμενο", ['<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><mc:Choice Requires="w14"><w:p><w:r><w:t>Νέα μορφή</w:t></w:r></w:p>'
                             '</mc:Choice><mc:Fallback><w:p><w:r><w:t>Νέα μορφή</w:t></w:r></w:p></mc:Fallback>'
                             '</mc:AlternateContent>'], None, False, "Κείμενο\n\nΝέα μορφή",
                 id="paragraph written twice for old word"),
    pytest.param("Κείμενο", ['<w:sdt><w:sdtPr><w:alias w:val="Σώμα"/></w:sdtPr><w:sdtContent>'
                             '<w:sdt><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Table of Contents"/></w:docPartObj>'
                             '</w:sdtPr><w:sdtContent><w:p><w:r><w:t>Εισαγωγή 1</w:t></w:r></w:p></w:sdtContent></w:sdt>'
                             '<w:p><w:r><w:t>Το σώμα της εργασίας</w:t></w:r></w:p></w:sdtContent></w:sdt>'], None, False,
                 "Κείμενο\n\nΤο σώμα της εργασίας", id="table of contents inside another content control"),
])
def test_read_docx_text_inside_other_elements(tmp_path, first, xml_parts, last, inside, expected):
    assert read_docx(docx_with_xml(tmp_path, first, xml_parts, last, inside)) == expected


# rows word repeats at the top of every page are the header; google writes w:val="0" on every row
@pytest.mark.parametrize("rows, marks, expected", [
    ([["2023", "2024"], ["120", "340"]], [], "2023 | 2024\n120 | 340"),
    ([["2023", "2024"], ["120", "340"]], [None], "2023: 120; 2024: 340"),
    ([["2023", "2024"], ["120", "340"]], ["0", "0"], "2023 | 2024\n120 | 340"),
    ([["2023", "2024"], ["120", "340"]], ["false"], "2023 | 2024\n120 | 340"),
    ([["Βαθμοί", "Βαθμοί"], ["Γραπτό", "Προφορικό"], ["9", "8"]], ["1", "true"],
                 "Βαθμοί - Γραπτό: 9; Βαθμοί - Προφορικό: 8"),
    pytest.param([["2023", "2024"], ["120", "340"], ["150", "360"]], ["0", None],
                 "2023 | 2024\n120 | 340\n150 | 360"),
], ids=["no mark", "marked", "google: val 0", "val false", "two marked rows", "only rows at the top count"])
def test_read_docx_table_header_rows(tmp_path, rows, marks, expected):
    d = docx.Document()
    table = d.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    for row, value in zip(table.rows, marks):
        val = "" if value is None else f' w:val="{value}"'
        row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader xmlns:w="{W_NS}"{val}/>'))
    f = tmp_path / "header.docx"
    d.save(str(f))

    assert read_docx(f) == expected


# the alt text of a picture: the description, else the title (like odt)
@pytest.mark.parametrize("descr, title, expected", [
    ("Διάγραμμα", "Λογότυπο", "[Image: Διάγραμμα]"),
    ("", "Λογότυπο", "[Image: Λογότυπο]"),
    ("", "", ""),
])
def test_read_docx_picture_alt_text(tmp_path, descr, title, expected):
    image = tmp_path / "img.png"
    Image.new("RGB", (60, 30), "white").save(image)
    d = docx.Document()
    d.add_picture(str(image), width=Inches(1))
    d.inline_shapes[-1]._inline.docPr.set("descr", descr)
    d.inline_shapes[-1]._inline.docPr.set("title", title)
    f = tmp_path / "picture.docx"
    d.save(str(f))

    assert read_docx(f) == expected


# a heading style inside a table cell is not a heading of the document
def test_read_docx_heading_inside_a_cell(tmp_path):
    d = docx.Document()
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text, table.cell(1, 0).text = "Ενότητα", "Σημείωση", "Α"
    table.cell(1, 1).paragraphs[0].text = "Τίτλος"
    table.cell(1, 1).paragraphs[0].style = "Heading 1"
    f = tmp_path / "cell.docx"
    d.save(str(f))

    assert read_docx(f) == "Ενότητα: Α; Σημείωση: Τίτλος"


# a line break (shift+enter) is a new line, and a slide title is always one line
def test_read_pptx_line_breaks(tmp_path):
    p = pptx.Presentation()
    slide = p.slides.add_slide(p.slide_layouts[LAYOUT_TITLE_ONLY])
    title = slide.shapes.title
    assert title is not None
    title.text = "Γραμμή 1"
    title.text_frame.paragraphs[0].add_line_break()
    title.text_frame.paragraphs[0].add_run().text = "Γραμμή 2"
    title.text_frame.add_paragraph().text = "  "
    title.text_frame.add_paragraph().text = "Γραμμή 3"
    for frame in (slide.shapes.add_textbox(PptxInches(1), PptxInches(3), PptxInches(3), PptxInches(1)).text_frame,
                  slide.notes_slide.notes_text_frame):
        assert frame is not None
        frame.text = "α"
        frame.paragraphs[0].add_line_break()
        frame.paragraphs[0].add_run().text = "β"
    table = slide.shapes.add_table(2, 2, PptxInches(1), PptxInches(5), PptxInches(4), PptxInches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text, table.cell(1, 0).text = "Όνομα", "Σχόλιο", "Μαρία"
    cell = table.cell(1, 1).text_frame
    cell.text = "α"
    cell.paragraphs[0].add_line_break()
    cell.paragraphs[0].add_run().text = "β"
    f = tmp_path / "breaks.pptx"
    p.save(str(f))

    assert read_pptx(f) == ("# Slide 1: Γραμμή 1 Γραμμή 2 Γραμμή 3\n\nα\nβ\n\n"
                            "Όνομα: Μαρία; Σχόλιο: α / β\n\nNotes: α\nβ")


# a hidden slide is read like the others
def test_read_pptx_hidden_slide(tmp_path):
    p = pptx.Presentation()
    for text in ("Α", "Β"):
        p.slides.add_slide(p.slide_layouts[LAYOUT_BLANK]).shapes.add_textbox(
            PptxInches(1), PptxInches(1), PptxInches(3), PptxInches(1)).text_frame.text = text
    p.slides[1]._element.set("show", "0")
    f = tmp_path / "hidden.pptx"
    p.save(str(f))

    assert read_pptx(f) == "# Slide 1\n\nΑ\n\n# Slide 2\n\nΒ"


# the title of a chart should be in the text
@pytest.mark.xfail(strict=True, reason="L33: charts are not read")
def test_read_pptx_chart_title(tmp_path):
    p = pptx.Presentation()
    slide = p.slides.add_slide(p.slide_layouts[LAYOUT_BLANK])
    data = CategoryChartData()
    data.categories = ["2023", "2024"]
    data.add_series("Φοιτητές", (120, 150))
    # wrong type hints in python-pptx
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, PptxInches(1), PptxInches(1),
                                   PptxInches(6), PptxInches(4), data)  # pyright: ignore[reportArgumentType]
    chart = frame.chart  # pyright: ignore[reportAttributeAccessIssue]
    chart.has_title = True
    chart.chart_title.text_frame.text = "Εγγραφές ανά έτος"
    f = tmp_path / "chart.pptx"
    p.save(str(f))

    assert "Εγγραφές ανά έτος" in read_pptx(f)


# writes an rtf string to a file
def write_rtf(tmp_path, content, name="doc.rtf", encoding="ascii"):
    f = tmp_path / name
    f.write_bytes(content.encode(encoding))
    return f


# rtf files
@pytest.mark.parametrize("rtf, encoding, expected", [
    pytest.param(RTF_GREEK_CP1253, "ascii", EXPECTED_RTF, id="Greek as cp1253 escapes"),
    pytest.param(r"{\rtf1\ansi\ansicpg1252\uc1 \u915?\u949?\u953?\u945? \u963?\u959?\u965?\par}",
                 "ascii", "Γεια σου", id="Greek as unicode escapes"),
    pytest.param("{\\rtf1\\ansi\\ansicpg1253 Γεια σου κόσμε\\par}", "cp1253", "Γεια σου κόσμε",
                 id="Greek written directly in cp1253"),
    pytest.param("\r\n  {\\rtf1\\ansi \\par\\par Text\\par\\par\\par}", "ascii", "Text",
                 id="blank lines and spaces stripped at both ends"),
    pytest.param(r"{\rtf1\ansi}", "ascii", "", id="empty document"),
    pytest.param(r"{\rtf1\ansi Before {\pict\pngblip\picw10\pich10 89504e470d0a1a0a0000}after\par}",
                 "ascii", "Before after", id="pictures dropped"),
    pytest.param(r"{\rtf1\ansi\trowd\cellx2000\cellx4000 Name\cell Grade\cell\row"
                 r"\trowd\cellx2000\cellx4000 Maria\cell 9\cell\row\pard After table\par}",
                 "ascii", "Name|Grade|\nMaria|9|\nAfter table", id="tables stay as cells with bars"),
    pytest.param(r"{\rtf1\ansi\ansicpg1253{\fonttbl{\f0 Arial;}}"
                 r"{\info{\title Secret title}{\author Some Author}}\f0 Body text.\par}",
                 "ascii", "Body text.", id="document properties are not body text"),
])
def test_read_rtf(tmp_path, rtf, encoding, expected):
    assert read_rtf(write_rtf(tmp_path, rtf, encoding=encoding)) == expected


# a footnote should not break the sentence
def test_read_rtf_footnote_is_not_glued_into_the_sentence(tmp_path):
    rtf = (r"{\rtf1\ansi{\header Page header}"
           r"Main sentence{\super\chftn{\footnote\pard\plain\chftn Footnote text.}} continues.\par}")

    assert "Main sentence continues." in read_rtf(write_rtf(tmp_path, rtf))


# a file that doesn't start with {\rtf1 is not rtf
@pytest.mark.parametrize("content", ["this is not an rtf file", "", "{\\rtfX not really}"])
def test_read_rtf_rejects_files_that_are_not_rtf(tmp_path, content):
    with pytest.raises(ValueError):
        read_rtf(write_rtf(tmp_path, content))


# ---- line breaks inside one paragraph (D52): the same text in docx, odt, pptx and odp ----
BR = None  # a line break (shift+enter) in the parts of a paragraph


def breaks_docx(parts, f):
    d = docx.Document()
    p = d.add_paragraph()
    for part in parts:
        if part is BR:
            p.add_run().add_break()
        else:
            p.add_run(part)
    d.save(str(f))
    return read_docx(f)


def odf_paragraph(parts):
    p = OdfP()
    for part in parts:
        if part is BR:
            p.addElement(OdfLineBreak())
        else:
            p.addText(part)
    return p


def breaks_odt(parts, f):
    doc = OpenDocumentText()
    doc.text.addElement(odf_paragraph(parts))  # pyright: ignore[reportAttributeAccessIssue] (odfpy adds it at runtime)
    doc.save(str(f))
    return read_odf(f)


def breaks_pptx(parts, f):
    p = pptx.Presentation()
    box = p.slides.add_slide(p.slide_layouts[LAYOUT_BLANK]).shapes.add_textbox(
        PptxInches(1), PptxInches(1), PptxInches(4), PptxInches(2))
    paragraph = box.text_frame.paragraphs[0]
    for part in parts:
        if part is BR:
            paragraph.add_line_break()
        else:
            paragraph.add_run().text = part
    p.save(str(f))
    return read_pptx(f).removeprefix("# Slide 1\n\n")


def breaks_odp(parts, f):
    doc = OpenDocumentPresentation()
    frame = Frame()
    box = OdfTextBox()
    box.addElement(odf_paragraph(parts))
    frame.addElement(box)
    page = Page(masterpagename="Default")
    page.addElement(frame)
    doc.presentation.addElement(page)  # pyright: ignore[reportAttributeAccessIssue] (odfpy adds it at runtime)
    doc.save(str(f))
    return read_odf(f).removeprefix("# Slide 1\n\n")


BREAK_FORMATS = {"docx": breaks_docx, "odt": breaks_odt, "pptx": breaks_pptx, "odp": breaks_odp}
BREAK_CASES = {
    "one break": (["α", BR, "β"], "α\nβ"),
    "two breaks": (["α", BR, BR, "β"], "α\n\nβ"),
    "three breaks": (["α", BR, BR, BR, "β"], "α\n\nβ"),
    "spaces between two breaks": (["α", BR, "  ", BR, "β"], "α\n\nβ"),
    "breaks at the start and the end": ([BR, BR, "α", BR, BR], "α"),
    "two blank lines between three lines": (["α", BR, BR, "β", BR, BR, BR, "γ"], "α\n\nβ\n\nγ"),
    "one break after an empty line": (["α", BR, BR, "β", BR, "γ"], "α\n\nβ\nγ"),
}


# one break is a new line, two or more are one empty line, in every format
@pytest.mark.parametrize("fmt, case", [pytest.param(fmt, case, id=f"{fmt}, {case}")
                                       for fmt, case in itertools.product(BREAK_FORMATS, BREAK_CASES)])
def test_line_breaks_in_every_format(tmp_path, fmt, case):
    parts, expected = BREAK_CASES[case]
    assert BREAK_FORMATS[fmt](parts, tmp_path / f"breaks.{fmt}") == expected


# an empty paragraph in a text box or in the notes is dropped, like in odp
def test_read_pptx_empty_paragraphs_are_dropped(tmp_path):
    p = pptx.Presentation()
    frame = p.slides.add_slide(p.slide_layouts[LAYOUT_BLANK]).shapes.add_textbox(
        PptxInches(1), PptxInches(1), PptxInches(4), PptxInches(2)).text_frame
    frame.text = "α"
    frame.add_paragraph().text = "   "
    frame.add_paragraph().text = "β"
    notes = p.slides[0].notes_slide.notes_text_frame
    assert notes is not None
    notes.text = "γ"
    notes.add_paragraph().text = ""
    notes.add_paragraph().text = "δ"
    f = tmp_path / "empty.pptx"
    p.save(str(f))

    assert read_pptx(f) == "# Slide 1\n\nα\nβ\n\nNotes: γ\nδ"


# ---- links keep their address (D50): the same text in docx, odt, pptx and odp ----
# a paragraph is a list of plain strings and links (pieces of the link text, address); "#..." is a place in the same file
def links_docx(parts, f):
    d = docx.Document()
    p = d.add_paragraph()
    for part in parts:
        if isinstance(part, str):
            p.add_run(part)
            continue
        pieces, url = part
        runs = "".join(f'<w:r><w:t xml:space="preserve">{piece}</w:t></w:r>' for piece in pieces)
        if url.startswith("#"):
            target = f'w:anchor="{url[1:]}"'
        else:
            target = f'r:id="{d.part.relate_to(url, DOCX_RT.HYPERLINK, is_external=True)}"'
        p._p.append(parse_xml(f'<w:hyperlink xmlns:w="{W_NS}" xmlns:r="{R_NS}" {target}>{runs}</w:hyperlink>'))
    d.save(str(f))
    return read_docx(f)


def odf_link_paragraph(parts):
    p = OdfP()
    for part in parts:
        if isinstance(part, str):
            p.addText(part)
            continue
        pieces, url = part
        link = OdfA(href=url, type="simple")
        for piece in pieces:
            link.addElement(OdfSpan(text=piece))
        p.addElement(link)
    return p


def links_odt(parts, f):
    doc = OpenDocumentText()
    doc.text.addElement(odf_link_paragraph(parts))  # pyright: ignore[reportAttributeAccessIssue] (odfpy adds it at runtime)
    doc.save(str(f))
    return read_odf(f)


def links_pptx(parts, f):
    p = pptx.Presentation()
    slide = p.slides.add_slide(p.slide_layouts[LAYOUT_BLANK])
    target = p.slides.add_slide(p.slide_layouts[LAYOUT_BLANK])
    paragraph = slide.shapes.add_textbox(PptxInches(1), PptxInches(1), PptxInches(6), PptxInches(1)).text_frame.paragraphs[0]
    for part in parts:
        if isinstance(part, str):
            paragraph.add_run().text = part
            continue
        pieces, url = part
        for piece in pieces:
            run = paragraph.add_run()
            run.text = piece
            run.hyperlink.address = "https://placeholder" if url.startswith("#") else url
            if url.startswith("#"):  # a jump to another slide
                link = run._r.find(pptx_qn("a:rPr")).find(pptx_qn("a:hlinkClick"))
                link.set(pptx_qn("r:id"), slide.part.relate_to(target.part, PPTX_RT.SLIDE))
                link.set("action", "ppaction://hlinksldjump")
    p.save(str(f))
    return read_pptx(f).removeprefix("# Slide 1\n\n").removesuffix("\n\n# Slide 2")


def links_odp(parts, f):
    doc = OpenDocumentPresentation()
    frame = Frame()
    box = OdfTextBox()
    box.addElement(odf_link_paragraph(parts))
    frame.addElement(box)
    page = Page(masterpagename="Default")
    page.addElement(frame)
    doc.presentation.addElement(page)  # pyright: ignore[reportAttributeAccessIssue] (odfpy adds it at runtime)
    doc.save(str(f))
    return read_odf(f).removeprefix("# Slide 1\n\n")


R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
LINK_FORMATS = {"docx": links_docx, "odt": links_odt, "pptx": links_pptx, "odp": links_odp}
LINK_CASES = {
    "text is not the address": (["Δες τον ", (["οδηγό"], "https://x.gr/guide"), " για λεπτομέρειες."],
                                "Δες τον οδηγό (https://x.gr/guide) για λεπτομέρειες."),
    "text is the address": ([(["https://elearn.uoc.gr"], "https://elearn.uoc.gr/")], "https://elearn.uoc.gr"),
    "address without http and www": ([(["csd.uoc.gr"], "http://www.csd.uoc.gr/")], "csd.uoc.gr"),
    "email shown as the address": ([(["maria@uoc.gr"], "mailto:maria@uoc.gr")], "maria@uoc.gr"),
    "email behind a name": (["Γράψε στη ", (["Μαρία"], "mailto:maria@uoc.gr")], "Γράψε στη Μαρία (mailto:maria@uoc.gr)"),
    "link to a place in the same file": (["Δες την ", (["Εισαγωγή"], "#intro")], "Δες την Εισαγωγή"),
    "space at the end of the link text": ([(["εδώ "], "https://x.gr"), "για"], "εδώ (https://x.gr) για"),
    "address written in capitals": ([(["ELEARN.UOC.GR"], "https://elearn.uoc.gr/")], "ELEARN.UOC.GR"),
    "link without text": (["Κείμενο", ([""], "https://x.gr")], "Κείμενο"),
    "space at the start of the link text": (["Δες τον", ([" οδηγό"], "https://x.gr/guide"), " για"],
                                            "Δες τον οδηγό (https://x.gr/guide) για"),
    "link text in two pieces": ([(["οδη", "γό"], "https://x.gr/guide")], "οδηγό (https://x.gr/guide)"),
    "two different links": ([(["α"], "https://a.gr"), " και ", (["β"], "https://b.gr")],
                            "α (https://a.gr) και β (https://b.gr)"),
}


# the address of a link is kept after its text, unless the text already is the address or it is inside the file
@pytest.mark.parametrize("fmt, case", [pytest.param(fmt, case, id=f"{fmt}, {case}")
                                       for fmt, case in itertools.product(LINK_FORMATS, LINK_CASES)])
def test_links_in_every_format(tmp_path, fmt, case):
    parts, expected = LINK_CASES[case]
    assert LINK_FORMATS[fmt](parts, tmp_path / f"links.{fmt}") == expected


# a field (like the slide number) is text of the paragraph
def test_read_pptx_field_in_a_paragraph(tmp_path):
    p = pptx.Presentation()
    paragraph = p.slides.add_slide(p.slide_layouts[LAYOUT_BLANK]).shapes.add_textbox(
        PptxInches(1), PptxInches(1), PptxInches(4), PptxInches(1)).text_frame.paragraphs[0]
    paragraph.add_run().text = "Διαφάνεια "
    paragraph._p.append(parse_xml(
        '<a:fld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}"'
        ' type="slidenum"><a:t>1</a:t></a:fld>'))
    f = tmp_path / "field.pptx"
    p.save(str(f))

    assert read_pptx(f) == "# Slide 1\n\nΔιαφάνεια 1"


# (file name, text in the file, text after parse_file)
NORMALIZE_CASES = {
    "micro sign in a Greek word": ("notes.txt", "Το σήµα", "Το σήμα"),
    "math letters in Greek words": ("notes.txt", "ϕως ϐάση ϑέμα ϱοή", "φως βάση θέμα ροή"),
    "tonos before every capital vowel": (
        "notes.txt", "΄Αλφα ΄Εψιλον ΄Ητα ΄Ιωτα ΄Ομικρον ΄Υψιλον ΄Ωμέγα", "Άλφα Έψιλον Ήτα Ίωτα Όμικρον Ύψιλον Ώμέγα"),
    "acute accent before a capital": ("notes.txt", "´Ασκηση", "Άσκηση"),
    "ligature and no-break space": ("notes.txt", "ﬁle name", "file name"),
    "table cells": ("grades.csv", "Λέξη;Βαθμός\nσήµα;9\n", "Λέξη: σήμα; Βαθμός: 9"),
    "markdown heading": ("notes.md", "# ΄Ασκηση 1\n\nΤο σήµα", "# Άσκηση 1\n\nΤο σήμα"),
}


# look-alike letters and a tonos written apart become the normal Greek letters
@pytest.mark.parametrize("name, written, expected", NORMALIZE_CASES.values(), ids=NORMALIZE_CASES.keys())
def test_parse_file_normalizes_unicode(tmp_path, name, written, expected):
    assert parse_one(tmp_path, name, written.encode()).raw_text == expected


# text that is already right does not change
@pytest.mark.parametrize("text", ["Άσκηση και μάθημα", "Ελληνικά, English, 2024-03-01 15:30", "ΆΣΚΗΣΗ ΜΆΘΗΜΑ"])
def test_parse_file_keeps_normal_text(tmp_path, text):
    assert parse_one(tmp_path, "notes.txt", text.encode()).raw_text == text


# code is kept as written, a string in a program must not change
def test_parse_file_leaves_code_as_written(tmp_path):
    code = 'MICRO = "µ"\nPHI = "ϕ"\nTONOS = "΄Α"\nFI = "ﬁ"\n'
    assert parse_one(tmp_path, "units.py", code.encode()).raw_text == code




# utf-8 text that a program read as latin-1 and saved again ("Α" -> "Î\x91")
def as_latin1(text):
    return text.encode("utf-8").decode("latin-1")


# utf-8 text that a program read as windows-1252 ("υ" -> "Ï…"); bytes that 1252 does not have stay as latin-1
def as_cp1252(text):
    return "".join(bytes([b]).decode("cp1252", errors="ignore") or chr(b) for b in text.encode("utf-8"))


# (file name, text written in it, text that comes out)
MOJIBAKE_CASES = {
    "greek read as latin-1": ("notes.txt", as_latin1("Αποτέλεσμα εικόνας για ρομπότ"), "Αποτέλεσμα εικόνας για ρομπότ"),
    "greek read as windows-1252": ("notes.txt", as_cp1252("Το αρχείο του μαθήματος"), "Το αρχείο του μαθήματος"),
    "french read as latin-1": ("notes.txt", as_latin1("Le café est très bon"), "Le café est très bon"),
    "euro sign and quotes of three bytes": ("notes.txt", as_cp1252("Κόστος 5€ “τελικό”"), "Κόστος 5€ “τελικό”"),
    "only the broken part changes": (
        "notes.txt", "Το σωστό κείμενο και " + as_latin1("το χαλασμένο") + " μαζί.", "Το σωστό κείμενο και το χαλασμένο μαζί."),
    "fixed text is normalized after": ("notes.txt", as_latin1("Το σήµα"), "Το σήμα"),
    "table cells": ("grades.csv", "Όνομα;Σχόλιο\nΜαρία;" + as_latin1("Πολύ καλή") + "\n", "Όνομα: Μαρία; Σχόλιο: Πολύ καλή"),
    "markdown heading": ("notes.md", "# " + as_latin1("Άσκηση") + "\n\nΚείμενο", "# Άσκηση\n\nΚείμενο"),
}


@pytest.mark.parametrize("name, written, expected", MOJIBAKE_CASES.values(), ids=MOJIBAKE_CASES.keys())
def test_parse_file_fixes_mojibake(tmp_path, name, written, expected):
    assert parse_one(tmp_path, name, written.encode()).raw_text == expected


# a word document with mojibake in a paragraph is fixed too
def test_parse_file_fixes_mojibake_in_docx(tmp_path):
    def make(f):
        d = docx.Document()
        d.add_paragraph(as_latin1("Αποτέλεσμα εικόνας"))
        d.save(f)
    assert parse_one(tmp_path, "notes.docx", built(tmp_path, "made.docx", make)).raw_text == "Αποτέλεσμα εικόνας"


# text with letters that only look like mojibake stays as it is
MOJIBAKE_KEPT = {
    "capitals with accents before ascii": "Île-de-France, São Paulo, Ãgua",
    "signs alone": "Τιμή: 5 € ή 6 £, θερμοκρασία 20 °C, © 2024",
    "letters in a row that are not utf-8": "ÀÁÂ ÃÄÅ",
    "a lead letter at the end": "Το τέλος είναι Î",
}


@pytest.mark.parametrize("text", MOJIBAKE_KEPT.values(), ids=MOJIBAKE_KEPT.keys())
def test_parse_file_keeps_text_that_is_not_mojibake(tmp_path, text):
    assert parse_one(tmp_path, "notes.txt", text.encode()).raw_text == text


# code is kept as written, even when it looks like mojibake
def test_parse_file_leaves_mojibake_in_code(tmp_path):
    code = 'BROKEN = "' + as_latin1("Α") + '"\n'
    assert parse_one(tmp_path, "broken.py", code.encode()).raw_text == code


# programs in the alpha language of the compiler course are code, so nothing in them is changed
@pytest.mark.parametrize("name", ["test1.asc", "queue.alpha", "MAIN.ASC"])
def test_alpha_files_are_code(tmp_path, name):
    program = 'x = "µ";\nprint("΄Α", x);\n'
    doc = parse_one(tmp_path, name, program.encode())
    assert (doc.content_type, doc.raw_text) == ("code", program)


# inputs and outputs of the compiler tests stay plain text
@pytest.mark.parametrize("name", ["test1.in", "test1.out"])
def test_compiler_inputs_and_outputs_are_text(tmp_path, name):
    assert parse_one(tmp_path, name, b"quad 1: assign x 5\n").content_type == "prose"
